"""
Reviewer-ready Health Factor (HF) calculator for the current İZSU dataset.

This module implements the manuscript's WAWQI-based HF formulation and can be
used either as a reusable library or as a command-line report generator.

Command-line example
--------------------
python hf_calculator.py --input ../data/izsu_data_cleaned.csv \
    --output-dir ../outputs/hf

Generated files
---------------
- izsu_health_factor_recalculated.csv
- hf_parameter_configuration.csv
- hf_validation_summary.csv
- HF_Technical_Validation_Report.docx
"""

from __future__ import annotations

import argparse
import math
import re
from pathlib import Path
from typing import Any, Mapping

import numpy as np
import pandas as pd

try:  # Package import: from utils.hf_calculator import ...
    from .parameters import (
        FAIL_FAST_HF,
        HF_CAUTION_MIN,
        HF_GOOD_MIN,
        HF_MAX,
        HF_MIN,
        PARAMETERS,
        UNIT_WEIGHTS,
        WAWQI_K,
        WAWQI_PARAMETER_NAMES,
        canonical_parameter_name,
    )
except ImportError:  # Direct execution: python utils/hf_calculator.py
    from parameters import (  # type: ignore
        FAIL_FAST_HF,
        HF_CAUTION_MIN,
        HF_GOOD_MIN,
        HF_MAX,
        HF_MIN,
        PARAMETERS,
        UNIT_WEIGHTS,
        WAWQI_K,
        WAWQI_PARAMETER_NAMES,
        canonical_parameter_name,
    )


ACCEPTABLE_GOOD_VALUES = {
    "uygun",
    "geçerli",
    "gecerli",
    "0",
    "yok",
    "yoktur",
    "nd",
    "-",
    "—",
}


def normalize_unit(unit: object) -> str:
    """Normalize unit spelling without silently changing its magnitude."""

    if unit is None or (isinstance(unit, float) and pd.isna(unit)):
        return ""
    text = re.sub(r"\s+", " ", str(unit).strip())
    text = text.replace("µ", "μ")
    text = re.sub(r"(?i)\bu[gG]/L\b", "μg/L", text)
    return text


def convert_unit(value: float, source_unit: object, target_unit: str | None) -> float:
    """Convert supported concentration units and reject unknown mismatches."""

    if pd.isna(value):
        return math.nan
    if target_unit is None:
        return float(value)

    source = normalize_unit(source_unit)
    target = normalize_unit(target_unit)

    # Missing source units are accepted only when the dataset's parameter
    # configuration unambiguously defines the expected unit.
    if not source or source == target:
        return float(value)
    if (source, target) == ("μg/L", "mg/L"):
        return float(value) / 1000.0
    if (source, target) == ("mg/L", "μg/L"):
        return float(value) * 1000.0

    return math.nan


def classify_hf(hf: float, fail_fast: bool = False) -> str:
    """Apply the manuscript thresholds exactly."""

    if fail_fast:
        return "Risk"
    if pd.isna(hf):
        return "Unknown"
    if hf >= HF_GOOD_MIN:
        return "Good"
    if hf >= HF_CAUTION_MIN:
        return "Caution"
    return "Risk"


def quality_rating_numeric(value: float, standard: float, ideal: float = 0.0) -> float:
    """WAWQI quality rating Qn for a positive numeric standard."""

    if pd.isna(value) or standard <= ideal:
        return math.nan
    return float(abs(value - ideal) / (standard - ideal) * 100.0)


def quality_rating_range(
    value: float,
    lower: float,
    upper: float,
    ideal: float,
) -> float:
    """Piecewise WAWQI quality rating for a range-based standard such as pH."""

    if pd.isna(value):
        return math.nan
    denominator = (upper - ideal) if value >= ideal else (ideal - lower)
    if denominator <= 0:
        return math.nan
    return float(abs(value - ideal) / denominator * 100.0)


def acceptable_score(raw_value: object) -> float:
    """Return 1.0 for explicit acceptable text, otherwise NaN.

    A numeric turbidity value is not assigned a regulatory score unless a
    positive numeric standard is provided by the source/configuration.
    """

    if raw_value is None or (isinstance(raw_value, float) and pd.isna(raw_value)):
        return math.nan
    text = str(raw_value).strip().casefold()
    return 1.0 if text in ACCEPTABLE_GOOD_VALUES else math.nan


def score_from_quality_rating(qn: float) -> float:
    """Convert risk-oriented Qn to a safety-oriented 0–1 parameter score."""

    if pd.isna(qn):
        return math.nan
    return float(np.clip(1.0 - max(qn, 0.0) / 100.0, 0.0, 1.0))


def _mean_numeric(series: pd.Series) -> float:
    values = pd.to_numeric(series, errors="coerce")
    return float(values.mean()) if values.notna().any() else math.nan


def _first_nonempty(series: pd.Series) -> str:
    values = series.dropna().astype(str).str.strip()
    values = values[values != ""]
    return "" if values.empty else values.iloc[0]


def calculate_parameter_result(
    parameter_name: object,
    value: float,
    unit: object = "",
    raw_value: object = None,
) -> dict[str, Any]:
    """Calculate Qn, safety score, and fail-fast state for one parameter."""

    canonical = canonical_parameter_name(parameter_name)
    spec = PARAMETERS.get(canonical)
    result: dict[str, Any] = {
        "parameter": canonical,
        "value": value,
        "unit": normalize_unit(unit),
        "standardized_value": math.nan,
        "quality_rating": math.nan,
        "score": math.nan,
        "fail_fast": False,
        "fail_fast_reason": "",
        "included_in_wawqi": False,
    }

    if spec is None:
        return result

    standardized = convert_unit(value, unit, spec.unit)
    result["standardized_value"] = standardized

    if spec.kind == "zero_standard":
        if not pd.isna(standardized):
            is_clean = float(standardized) == 0.0
            result["quality_rating"] = 0.0 if is_clean else 100.0
            result["score"] = 1.0 if is_clean else 0.0
            if spec.fail_fast and not is_clean:
                result["fail_fast"] = True
                result["fail_fast_reason"] = (
                    f"{canonical}={standardized:g} exceeds zero standard"
                )
        return result

    if spec.kind == "numeric":
        if pd.isna(standardized):
            return result
        standard = float(spec.standard)  # type: ignore[arg-type]
        ideal = float(spec.ideal or 0.0)
        qn = quality_rating_numeric(standardized, standard, ideal)
        result["quality_rating"] = qn
        result["score"] = score_from_quality_rating(qn)
        result["included_in_wawqi"] = spec.include_in_wawqi
        if spec.fail_fast and standardized > standard:
            result["fail_fast"] = True
            result["fail_fast_reason"] = (
                f"{canonical}={standardized:g} {spec.unit or ''} "
                f"exceeds Sn={standard:g}"
            ).strip()
        return result

    if spec.kind == "range":
        if pd.isna(standardized):
            return result
        lower, upper = spec.standard  # type: ignore[misc]
        ideal = float(spec.ideal)
        qn = quality_rating_range(standardized, float(lower), float(upper), ideal)
        result["quality_rating"] = qn
        result["score"] = score_from_quality_rating(qn)
        result["included_in_wawqi"] = spec.include_in_wawqi
        return result

    if spec.kind == "acceptable":
        score = acceptable_score(raw_value)
        result["score"] = score
        if not pd.isna(score):
            result["quality_rating"] = 0.0 if score == 1.0 else 100.0
        return result

    # Descriptive parameters are retained in the feature dataset but do not
    # enter the HF because no numeric standard is defined in the source table.
    return result


def calculate_hf_from_parameter_results(
    parameter_results: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    """Aggregate parameter results into WAWQI and HF."""

    fail_reasons = [
        str(item.get("fail_fast_reason", ""))
        for item in parameter_results.values()
        if bool(item.get("fail_fast", False))
    ]
    fail_reasons = [reason for reason in fail_reasons if reason]

    numerator = 0.0
    denominator = 0.0
    used_parameters: list[str] = []

    for name in WAWQI_PARAMETER_NAMES:
        item = parameter_results.get(name)
        if not item:
            continue
        qn = item.get("quality_rating", math.nan)
        if pd.isna(qn):
            continue
        weight = UNIT_WEIGHTS[name]
        numerator += weight * float(qn)
        denominator += weight
        used_parameters.append(name)

    wawqi = math.nan if denominator == 0.0 else numerator / denominator
    coverage = denominator  # Full configured WAWQI weights sum to 1.0.

    if fail_reasons:
        hf = FAIL_FAST_HF
        fail_fast = True
    elif pd.isna(wawqi):
        hf = math.nan
        fail_fast = False
    else:
        hf = float(np.clip(100.0 - wawqi, HF_MIN, HF_MAX))
        fail_fast = False

    scores = {
        f"{name}_score": parameter_results.get(name, {}).get("score", math.nan)
        for name in PARAMETERS
    }

    return {
        "HealthFactor": hf,
        "WAWQI": wawqi,
        "FailFast": fail_fast,
        "FailFastReason": "; ".join(fail_reasons),
        "RiskClass": classify_hf(hf, fail_fast),
        "WAWQICoverage": coverage,
        "WAWQIParametersUsed": len(used_parameters),
        "WAWQIParameterList": ", ".join(used_parameters),
        **scores,
    }


def calculate_hf_for_group(
    group_df: pd.DataFrame,
    *,
    parameter_col: str = "ParametreAdi",
    value_col: str = "Deger",
    unit_col: str = "Birim",
    raw_value_col: str = "DegerRaw",
) -> dict[str, Any]:
    """Calculate HF for one date-location group in long format."""

    required = {parameter_col, value_col}
    missing = required - set(group_df.columns)
    if missing:
        raise ValueError(f"HF group is missing columns: {sorted(missing)}")

    parameter_results: dict[str, dict[str, Any]] = {}

    for raw_name, sub in group_df.groupby(parameter_col, dropna=False):
        canonical = canonical_parameter_name(raw_name)
        value = _mean_numeric(sub[value_col])
        unit = _first_nonempty(sub[unit_col]) if unit_col in sub.columns else ""
        raw_value = (
            _first_nonempty(sub[raw_value_col])
            if raw_value_col in sub.columns
            else ""
        )
        parameter_results[canonical] = calculate_parameter_result(
            canonical,
            value,
            unit,
            raw_value,
        )

    return calculate_hf_from_parameter_results(parameter_results)


def calculate_hf_from_wide_row(row: Mapping[str, Any]) -> dict[str, Any]:
    """Calculate HF from a wide row containing canonical parameter columns."""

    parameter_results: dict[str, dict[str, Any]] = {}
    for canonical, spec in PARAMETERS.items():
        candidates = (canonical, *spec.aliases)
        found_name = next((name for name in candidates if name in row), None)
        if found_name is None:
            continue
        value = pd.to_numeric(pd.Series([row.get(found_name)]), errors="coerce").iloc[0]
        raw_value = row.get(f"{found_name}_Raw", row.get(found_name))
        unit = row.get(f"{found_name}_Unit", spec.unit or "")
        parameter_results[canonical] = calculate_parameter_result(
            canonical,
            float(value) if not pd.isna(value) else math.nan,
            unit,
            raw_value,
        )

    return calculate_hf_from_parameter_results(parameter_results)


def parameter_configuration_table() -> pd.DataFrame:
    """Return the exact parameter/standard/weight table used by the code."""

    rows = []
    for name, spec in PARAMETERS.items():
        if isinstance(spec.standard, tuple):
            standard = f"{spec.standard[0]}–{spec.standard[1]}"
        elif spec.standard is None:
            standard = "Not numerically defined"
        else:
            standard = spec.standard

        rows.append(
            {
                "Parameter": name,
                "Unit": spec.unit or "-",
                "Standard_Sn": standard,
                "Ideal_Vi": "-" if spec.ideal is None else spec.ideal,
                "Type": spec.kind,
                "Category": spec.category,
                "FailFast": spec.fail_fast,
                "IncludedInWAWQI": spec.include_in_wawqi,
                "UnitWeight_Wn": UNIT_WEIGHTS.get(name, math.nan),
            }
        )
    return pd.DataFrame(rows)


def validate_formula_invariants() -> pd.DataFrame:
    """Run deterministic checks required by the reviewer request."""

    checks: list[dict[str, Any]] = []

    checks.append(
        {
            "Check": "Configured WAWQI unit weights sum to one",
            "Expected": 1.0,
            "Observed": sum(UNIT_WEIGHTS.values()),
            "Passed": math.isclose(sum(UNIT_WEIGHTS.values()), 1.0, abs_tol=1e-12),
        }
    )

    for hf, expected in [(85.0, "Good"), (84.999, "Caution"), (60.0, "Caution"), (59.999, "Risk")]:
        observed = classify_hf(hf)
        checks.append(
            {
                "Check": f"Classification threshold at HF={hf}",
                "Expected": expected,
                "Observed": observed,
                "Passed": observed == expected,
            }
        )

    ideal_results = {}
    for name in WAWQI_PARAMETER_NAMES:
        spec = PARAMETERS[name]
        ideal = float(spec.ideal or 0.0)
        ideal_results[name] = calculate_parameter_result(name, ideal, spec.unit or "", str(ideal))
    ideal_hf = calculate_hf_from_parameter_results(ideal_results)
    checks.append(
        {
            "Check": "All numeric parameters at ideal values produce HF=100",
            "Expected": 100.0,
            "Observed": ideal_hf["HealthFactor"],
            "Passed": math.isclose(ideal_hf["HealthFactor"], 100.0, abs_tol=1e-9),
        }
    )

    e_coli_fail = calculate_hf_from_parameter_results(
        {"E.coli": calculate_parameter_result("E.coli", 1.0, "Sayı/100 ml", "1")}
    )
    checks.append(
        {
            "Check": "Zero-standard E.coli violation avoids division and triggers Risk",
            "Expected": "HF=0, Risk",
            "Observed": f"HF={e_coli_fail['HealthFactor']}, {e_coli_fail['RiskClass']}",
            "Passed": e_coli_fail["HealthFactor"] == 0.0 and e_coli_fail["RiskClass"] == "Risk",
        }
    )

    arsenic_fail = calculate_hf_from_parameter_results(
        {"Arsenik": calculate_parameter_result("Arsenik", 11.0, "μg/L", "11")}
    )
    checks.append(
        {
            "Check": "Arsenic above Sn triggers immediate Risk",
            "Expected": "HF=0, Risk",
            "Observed": f"HF={arsenic_fail['HealthFactor']}, {arsenic_fail['RiskClass']}",
            "Passed": arsenic_fail["HealthFactor"] == 0.0 and arsenic_fail["RiskClass"] == "Risk",
        }
    )

    return pd.DataFrame(checks)


def recalculate_hf_dataset(cleaned_df: pd.DataFrame) -> pd.DataFrame:
    """Recalculate HF from the current cleaned long-format dataset."""

    date_col = "Tarih_Clean" if "Tarih_Clean" in cleaned_df.columns else "Tarih"
    location_col = (
        "NoktaAdi_Clean" if "NoktaAdi_Clean" in cleaned_df.columns else "NoktaAdi"
    )
    parameter_col = (
        "ParametreAdi_Clean"
        if "ParametreAdi_Clean" in cleaned_df.columns
        else "ParametreAdi"
    )
    unit_col = "Birim_Clean" if "Birim_Clean" in cleaned_df.columns else "Birim"
    value_col = "Deger_Num" if "Deger_Num" in cleaned_df.columns else "Deger"
    raw_value_col = "DegerRaw" if "DegerRaw" in cleaned_df.columns else "Deger"

    required = {date_col, location_col, parameter_col, value_col}
    missing = required - set(cleaned_df.columns)
    if missing:
        raise ValueError(f"Cleaned dataset is missing columns: {sorted(missing)}")

    work = cleaned_df.copy()
    work[date_col] = pd.to_datetime(work[date_col], errors="coerce")
    work = work.dropna(subset=[date_col, location_col, parameter_col])

    rows: list[dict[str, Any]] = []
    for (date_value, location), group in work.groupby([date_col, location_col], sort=True):
        result = calculate_hf_for_group(
            group,
            parameter_col=parameter_col,
            value_col=value_col,
            unit_col=unit_col,
            raw_value_col=raw_value_col,
        )
        rows.append(
            {
                "Tarih": pd.Timestamp(date_value).date().isoformat(),
                "NoktaAdi": location,
                **result,
            }
        )

    return pd.DataFrame(rows).sort_values(["Tarih", "NoktaAdi"]).reset_index(drop=True)


def validation_summary_table(
    cleaned_df: pd.DataFrame,
    hf_df: pd.DataFrame,
    formula_checks: pd.DataFrame,
) -> pd.DataFrame:
    """Create a compact reviewer-facing HF validation summary."""

    class_counts = hf_df["RiskClass"].value_counts(dropna=False).to_dict()
    date_series = pd.to_datetime(hf_df["Tarih"], errors="coerce")

    rows = [
        ("Cleaned parameter rows", len(cleaned_df)),
        ("HF date-location observations", len(hf_df)),
        ("HF study start", date_series.min().date().isoformat()),
        ("HF study end", date_series.max().date().isoformat()),
        ("Unique locations", hf_df["NoktaAdi"].nunique()),
        ("Minimum HF", hf_df["HealthFactor"].min()),
        ("Mean HF", hf_df["HealthFactor"].mean()),
        ("Maximum HF", hf_df["HealthFactor"].max()),
        ("Good observations", class_counts.get("Good", 0)),
        ("Caution observations", class_counts.get("Caution", 0)),
        ("Risk observations", class_counts.get("Risk", 0)),
        ("Unknown observations", class_counts.get("Unknown", 0)),
        ("Fail-fast observations", int(hf_df["FailFast"].sum())),
        ("WAWQI proportionality constant K", WAWQI_K),
        ("WAWQI parameter count", len(WAWQI_PARAMETER_NAMES)),
        ("Formula checks passed", int(formula_checks["Passed"].sum())),
        ("Formula checks total", len(formula_checks)),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def write_docx_report(
    output_path: Path,
    input_path: Path,
    parameter_table: pd.DataFrame,
    summary_table: pd.DataFrame,
    formula_checks: pd.DataFrame,
) -> None:
    """Generate a Word report covering reviewer item E."""

    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT, WD_SECTION
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required to generate the report") from exc

    def set_cell_font(cell, size: float = 8.0, bold: bool = False) -> None:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(size)
                run.bold = bold

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Health Factor (HF) and WAWQI Technical Validation Report")
    run.bold = True
    run.font.size = Pt(16)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Input: {input_path.name}").italic = True

    document.add_heading("1. Purpose and implementation scope", level=1)
    document.add_paragraph(
        "This report documents the exact code path used to calculate the Health "
        "Factor (HF), the standards and ideal values used by the implementation, "
        "the treatment of zero-standard microbiological parameters, the direction "
        "of the HF scale, and the Good/Caution/Risk thresholds. HF is implemented "
        "as a safety-oriented transformation of WAWQI rather than as an unrelated "
        "new index."
    )

    document.add_heading("2. Implemented WAWQI and HF equations", level=1)
    for equation in [
        "K = 1 / Σ(1 / Sₙ)",
        "Wₙ = K / Sₙ",
        "Qₙ = |Vₙ - Vᵢ| / (Sₙ - Vᵢ) × 100",
        "WAWQI = Σ(Wₙ × Qₙ) / Σ(Wₙ available)",
        "HF = max(0, min(100, 100 - WAWQI))",
    ]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(equation).bold = True

    document.add_paragraph(
        "Because HF equals 100 minus WAWQI, an increase in any included quality "
        "rating Qₙ cannot increase HF. Therefore, larger HF values mathematically "
        "represent safer water, while larger WAWQI values represent poorer water "
        "quality."
    )

    document.add_heading("3. Zero-standard and fail-fast parameters", level=1)
    document.add_paragraph(
        "E. coli, coliform bacteria, and C. perfringens have Sₙ = 0. They are "
        "excluded from K/Sₙ unit-weight calculation, so the implementation never "
        "divides by zero. A measured value of zero receives a parameter safety "
        "score of 1. Any value greater than zero triggers the fail-fast rule, sets "
        "HF to 0, and classifies the observation as Risk. The same fail-fast policy "
        "is applied when arsenic or nitrite exceeds its positive standard limit, "
        "consistent with the manuscript's stated anti-eclipsing rule."
    )

    # Landscape section for the wide parameter configuration table.
    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Inches(0.45)
    landscape.bottom_margin = Inches(0.45)
    landscape.left_margin = Inches(0.45)
    landscape.right_margin = Inches(0.45)

    document.add_heading("4. Parameter configuration used by the code", level=1)
    display_columns = [
        "Parameter", "Unit", "Standard_Sn", "Ideal_Vi", "Type",
        "Category", "FailFast", "IncludedInWAWQI", "UnitWeight_Wn",
    ]
    headers = ["Parameter", "Unit", "Sₙ", "Vᵢ", "Type", "Category", "Fail-fast", "WAWQI", "Wₙ"]
    table = document.add_table(rows=1, cols=len(display_columns))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
        set_cell_font(table.rows[0].cells[index], 8.0, True)

    for _, row in parameter_table[display_columns].iterrows():
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if pd.isna(value):
                text = "-"
            elif isinstance(value, bool):
                text = "Yes" if value else "No"
            elif isinstance(value, float):
                if display_columns[index] == "UnitWeight_Wn":
                    text = f"{value:.6f}"
                else:
                    text = f"{value:g}"
            else:
                text = str(value).replace("Not numerically defined", "Not defined")
            cells[index].text = text
            set_cell_font(cells[index], 7.5)

    # Return to portrait for results/checks.
    portrait = document.add_section(WD_SECTION.NEW_PAGE)
    portrait.orientation = WD_ORIENT.PORTRAIT
    portrait.page_width, portrait.page_height = portrait.page_height, portrait.page_width
    portrait.top_margin = Inches(0.65)
    portrait.bottom_margin = Inches(0.65)
    portrait.left_margin = Inches(0.75)
    portrait.right_margin = Inches(0.75)

    document.add_heading("5. Dataset-level validation results", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    set_cell_font(table.rows[0].cells[0], 9.0, True)
    set_cell_font(table.rows[0].cells[1], 9.0, True)
    for _, row in summary_table.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["Metric"])
        value = row["Value"]
        cells[1].text = f"{value:.6f}" if isinstance(value, float) else str(value)
        set_cell_font(cells[0], 8.5)
        set_cell_font(cells[1], 8.5)

    document.add_heading("6. Deterministic formula and threshold checks", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, column in enumerate(["Check", "Expected", "Observed", "Passed"]):
        table.rows[0].cells[index].text = column
        set_cell_font(table.rows[0].cells[index], 8.5, True)
    for _, row in formula_checks.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["Check"])
        cells[1].text = str(row["Expected"])
        cells[2].text = str(row["Observed"])
        cells[3].text = "Yes" if bool(row["Passed"]) else "No"
        for cell in cells:
            set_cell_font(cell, 8.0)

    document.add_heading("7. Classification thresholds", level=1)
    document.add_paragraph("Good: HF ≥ 85", style="List Bullet")
    document.add_paragraph("Caution: 60 ≤ HF < 85", style="List Bullet")
    document.add_paragraph("Risk: HF < 60 or any fail-fast violation", style="List Bullet")

    document.add_heading("8. Reproducibility outputs", level=1)
    document.add_paragraph(
        "The same execution also writes the recalculated observation-level HF "
        "dataset, the parameter configuration table, deterministic formula checks, "
        "and a compact validation summary as CSV files. These files are intended "
        "to be consumed later by the classification, regression, table-generation, "
        "and final report scripts."
    )

    document.save(output_path)

def run_cli(input_path: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    cleaned = pd.read_csv(input_path, encoding="utf-8-sig")
    hf_df = recalculate_hf_dataset(cleaned)
    parameter_table = parameter_configuration_table()
    formula_checks = validate_formula_invariants()
    summary = validation_summary_table(cleaned, hf_df, formula_checks)

    hf_path = output_dir / "izsu_health_factor_recalculated.csv"
    parameter_path = output_dir / "hf_parameter_configuration.csv"
    checks_path = output_dir / "hf_formula_checks.csv"
    summary_path = output_dir / "hf_validation_summary.csv"
    report_path = output_dir / "HF_Technical_Validation_Report.docx"

    hf_df.to_csv(hf_path, index=False, encoding="utf-8-sig")
    parameter_table.to_csv(parameter_path, index=False, encoding="utf-8-sig")
    formula_checks.to_csv(checks_path, index=False, encoding="utf-8-sig")
    summary.to_csv(summary_path, index=False, encoding="utf-8-sig")
    write_docx_report(report_path, input_path, parameter_table, summary, formula_checks)

    print("--------------------------------------------------")
    print("HF / WAWQI validation completed")
    print(f"Input rows              : {len(cleaned)}")
    print(f"HF observations         : {len(hf_df)}")
    print(f"Class distribution      : {hf_df['RiskClass'].value_counts().to_dict()}")
    print(f"Fail-fast observations  : {int(hf_df['FailFast'].sum())}")
    print(f"Formula checks          : {int(formula_checks['Passed'].sum())}/{len(formula_checks)} passed")
    print(f"HF dataset              : {hf_path}")
    print(f"Parameter table         : {parameter_path}")
    print(f"Formula checks          : {checks_path}")
    print(f"Validation summary      : {summary_path}")
    print(f"Word report             : {report_path}")
    print("--------------------------------------------------")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Recalculate and validate WAWQI-based HF")
    parser.add_argument("--input", required=True, type=Path, help="Path to izsu_data_cleaned.csv")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("outputs/hf"),
        help="Directory for HF outputs",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    run_cli(args.input, args.output_dir)
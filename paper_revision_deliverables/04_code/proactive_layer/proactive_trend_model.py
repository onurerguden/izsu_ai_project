"""
proactive_trend_model.py

Reviewer-oriented Proactive Layer for the current İZSU dataset.

The model performs trend classification, not fixed-calendar forecasting.

Target definition
-----------------
For each physical sampling point i and observation time t:

    Delta_HF(i,t) = HF(i,t+1) - HF(i,t)

    Drop   if Delta_HF < -threshold
    Stable if -threshold <= Delta_HF <= threshold
    Rise   if Delta_HF > threshold

Here t+1 means the next AVAILABLE observation for the same sampling point.
Because measurement intervals are irregular, the actual number of days is
saved as HorizonDays and reported separately. The final observation of every
sampling point has no t+1 target and is excluded from supervised modeling.

Leakage-safe order
------------------
1. Load current izsu_features.csv.
2. Sort by NoktaId and date.
3. Create next-observation target and past-only features.
4. Split by unique dates into train/validation/test.
5. Fit imputer and scaler inside the training pipeline only.
6. Select the best model using validation Macro F1, MCC, minimum of Drop/Rise
   recall, and Balanced Accuracy.
7. Evaluate the selected and comparison models once on the untouched test set.

Required models
---------------
Extra Trees, Gradient Boosting, XGBoost, Random Forest, SVM, KNN.

Outputs
-------
- proactive_model_metrics.csv
- proactive_classwise_metrics.csv
- proactive_predictions.csv
- proactive_cv_metrics.csv
- proactive_split_summary.csv
- trend_label_summary.csv
- trend_horizon_summary.csv
- trend_horizon_distribution.csv
- trend_threshold_sensitivity.csv
- proactive_feature_list.csv
- data_leakage_audit.csv
- legacy_result_audit.csv
- figure_manifest.csv
- best_proactive_model.joblib
- Proactive_Layer_Technical_Report.docx
- figures/*.png at 300 DPI
"""

from __future__ import annotations

import argparse
import json
import math
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import joblib
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from sklearn.base import clone
from sklearn.ensemble import (
    ExtraTreesClassifier,
    GradientBoostingClassifier,
    RandomForestClassifier,
)
from sklearn.impute import SimpleImputer
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    confusion_matrix,
    f1_score,
    matthews_corrcoef,
    precision_recall_fscore_support,
    precision_score,
    recall_score,
)
from sklearn.neighbors import KNeighborsClassifier
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVC
from sklearn.utils.class_weight import compute_sample_weight

warnings.filterwarnings("ignore")

try:
    from xgboost import XGBClassifier

    XGBOOST_AVAILABLE = True
except Exception:
    XGBOOST_AVAILABLE = False


DATE_COL = "Tarih"
LOCATION_ID_COL = "NoktaId"
LOCATION_NAME_COL = "NoktaAdi"
HF_COL = "HealthFactor"
NEXT_HF_COL = "NextHealthFactor"
NEXT_DATE_COL = "NextObservationDate"
DELTA_COL = "FutureDeltaHF"
HORIZON_COL = "HorizonDays"
TARGET_COL = "TrendLabel"

CLASS_LABELS = ["Drop", "Stable", "Rise"]
CLASS_TO_INT = {"Drop": 0, "Stable": 1, "Rise": 2}
INT_TO_CLASS = {value: key for key, value in CLASS_TO_INT.items()}
RANDOM_STATE = 42

RAW_PARAMETER_CANDIDATES = [
    "Alüminyum",
    "Amonyum",
    "Arsenik",
    "Bulanıklık",
    "C.Perfringens",
    "Demir",
    "E.coli",
    "Klorür",
    "Koku",
    "Koliform Bakteri",
    "Nitrit",
    "Oksitlenebilirlik",
    "Renk",
    "Toplam Sertlik",
    "Tuzluluk",
    "pH",
    "İletkenlik",
]

FORBIDDEN_FEATURE_TERMS = {
    NEXT_HF_COL,
    NEXT_DATE_COL,
    DELTA_COL,
    HORIZON_COL,
    TARGET_COL,
    "RiskClass",
    "WAWQI",
    "FailFast",
    "FailFastReason",
}


@dataclass(frozen=True)
class DataSplit:
    train: pd.DataFrame
    validation: pd.DataFrame
    test: pd.DataFrame
    train_dates: list[pd.Timestamp]
    validation_dates: list[pd.Timestamp]
    test_dates: list[pd.Timestamp]


def find_input_csv(requested: str | None) -> Path:
    if requested:
        path = Path(requested)
        if path.exists():
            return path
        raise FileNotFoundError(f"Girdi CSV bulunamadı: {path}")

    script_dir = Path(__file__).resolve().parent
    candidates = [
        Path("data/data/izsu_features.csv"),
        Path("data/izsu_features.csv"),
        Path("izsu_features.csv"),
        script_dir / "data" / "izsu_features.csv",
        script_dir.parent / "data" / "izsu_features.csv",
        script_dir.parent / "data" / "data" / "izsu_features.csv",
    ]
    for path in candidates:
        if path.exists():
            return path

    raise FileNotFoundError(
        "izsu_features.csv bulunamadı. Dosya yolunu --input ile belirtin."
    )


def load_data(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, encoding="utf-8-sig")

    required = {DATE_COL, LOCATION_NAME_COL, HF_COL}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Eksik zorunlu sütunlar: {sorted(missing)}")

    if LOCATION_ID_COL not in df.columns:
        df[LOCATION_ID_COL] = df[LOCATION_NAME_COL].astype(str)

    df[DATE_COL] = pd.to_datetime(df[DATE_COL], errors="coerce")
    df[HF_COL] = pd.to_numeric(df[HF_COL], errors="coerce")
    df = df.dropna(
        subset=[DATE_COL, LOCATION_ID_COL, LOCATION_NAME_COL, HF_COL]
    ).copy()

    # Remove empty legacy spelling; canonical E.coli is used.
    if "E.Coli" in df.columns and df["E.Coli"].isna().all():
        df = df.drop(columns=["E.Coli"])

    duplicate_count = int(
        df.duplicated(subset=[DATE_COL, LOCATION_ID_COL]).sum()
    )
    if duplicate_count:
        raise ValueError(
            f"Tarih + NoktaId anahtarında {duplicate_count} duplicate bulundu."
        )

    numeric_candidates = set(RAW_PARAMETER_CANDIDATES) | {
        "Enlem",
        "Boylam",
    }
    for column in numeric_candidates:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")

    return df.sort_values(
        [LOCATION_ID_COL, DATE_COL], kind="stable"
    ).reset_index(drop=True)


def create_trend_target(
    df: pd.DataFrame,
    threshold: float,
) -> pd.DataFrame:
    if threshold <= 0:
        raise ValueError("Trend eşik değeri sıfırdan büyük olmalıdır.")

    result = df.copy()
    grouped = result.groupby(LOCATION_ID_COL, group_keys=False)

    result[NEXT_HF_COL] = grouped[HF_COL].shift(-1)
    result[NEXT_DATE_COL] = grouped[DATE_COL].shift(-1)
    result[DELTA_COL] = result[NEXT_HF_COL] - result[HF_COL]
    result[HORIZON_COL] = (
        result[NEXT_DATE_COL] - result[DATE_COL]
    ).dt.days

    result[TARGET_COL] = np.select(
        [
            result[DELTA_COL] < -threshold,
            result[DELTA_COL] > threshold,
        ],
        ["Drop", "Rise"],
        default="Stable",
    )

    # Last row of each location has no next observation and no valid label.
    result = result.dropna(
        subset=[NEXT_HF_COL, NEXT_DATE_COL, DELTA_COL, HORIZON_COL]
    ).copy()

    invalid_horizon = int((result[HORIZON_COL] <= 0).sum())
    if invalid_horizon:
        raise ValueError(
            f"{invalid_horizon} satırda pozitif olmayan tahmin ufku var."
        )

    return result.reset_index(drop=True)


def create_past_only_features(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    result = df.copy()
    group = result.groupby(LOCATION_ID_COL, group_keys=False)

    result["HF_Current"] = result[HF_COL]
    result["HF_Lag1"] = group[HF_COL].shift(1)
    result["HF_Lag2"] = group[HF_COL].shift(2)
    result["HF_Lag3"] = group[HF_COL].shift(3)

    result["HF_PastMean3"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(3, min_periods=1).mean()
    )
    result["HF_PastStd3"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(3, min_periods=2).std()
    )
    result["HF_PastMean7"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(7, min_periods=1).mean()
    )
    result["HF_PastStd7"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(7, min_periods=2).std()
    )

    result["HF_ChangeFromPrevious"] = (
        result["HF_Current"] - result["HF_Lag1"]
    )
    result["HF_PreviousChange"] = result["HF_Lag1"] - result["HF_Lag2"]
    result["HF_DifferenceFromPastMean3"] = (
        result["HF_Current"] - result["HF_PastMean3"]
    )

    previous_date = group[DATE_COL].shift(1)
    result["PreviousGapDays"] = (
        result[DATE_COL] - previous_date
    ).dt.days

    result["MonthSin"] = np.sin(
        2 * np.pi * result[DATE_COL].dt.month / 12.0
    )
    result["MonthCos"] = np.cos(
        2 * np.pi * result[DATE_COL].dt.month / 12.0
    )

    feature_columns = [
        column
        for column in RAW_PARAMETER_CANDIDATES
        if column in result.columns and not result[column].isna().all()
    ]

    feature_columns += [
        "HF_Current",
        "HF_Lag1",
        "HF_Lag2",
        "HF_Lag3",
        "HF_PastMean3",
        "HF_PastStd3",
        "HF_PastMean7",
        "HF_PastStd7",
        "HF_ChangeFromPrevious",
        "HF_PreviousChange",
        "HF_DifferenceFromPastMean3",
        "PreviousGapDays",
        "MonthSin",
        "MonthCos",
    ]

    for metadata_feature in ["Enlem", "Boylam"]:
        if (
            metadata_feature in result.columns
            and not result[metadata_feature].isna().all()
        ):
            feature_columns.append(metadata_feature)

    feature_columns = list(dict.fromkeys(feature_columns))

    forbidden = [
        column
        for column in feature_columns
        if column in FORBIDDEN_FEATURE_TERMS
        or column.endswith("_score")
    ]
    if forbidden:
        raise RuntimeError(
            f"Target veya gelecek bilgisi feature listesine sızdı: {forbidden}"
        )

    return result, feature_columns


def split_by_unique_dates(
    df: pd.DataFrame,
    train_ratio: float,
    validation_ratio: float,
) -> DataSplit:
    dates = sorted(pd.Series(df[DATE_COL].unique()).tolist())
    if len(dates) < 8:
        raise ValueError(
            "Train/validation/test ayrımı için en az 8 benzersiz tarih gerekir."
        )

    train_end = max(1, int(len(dates) * train_ratio))
    validation_end = max(
        train_end + 1,
        int(len(dates) * (train_ratio + validation_ratio)),
    )
    validation_end = min(validation_end, len(dates) - 1)

    train_dates = dates[:train_end]
    validation_dates = dates[train_end:validation_end]
    test_dates = dates[validation_end:]

    train = df[df[DATE_COL].isin(train_dates)].copy()
    validation = df[df[DATE_COL].isin(validation_dates)].copy()
    test = df[df[DATE_COL].isin(test_dates)].copy()

    for name, frame in [
        ("train", train),
        ("validation", validation),
        ("test", test),
    ]:
        if frame.empty:
            raise ValueError(f"{name} bölümü boş oluştu.")
        missing_classes = set(CLASS_LABELS) - set(frame[TARGET_COL].unique())
        if missing_classes:
            print(
                f"[UYARI] {name} bölümünde olmayan sınıflar: "
                f"{sorted(missing_classes)}"
            )

    return DataSplit(
        train=train,
        validation=validation,
        test=test,
        train_dates=train_dates,
        validation_dates=validation_dates,
        test_dates=test_dates,
    )


def get_models() -> dict[str, Any]:
    models: dict[str, Any] = {
        "Extra Trees": ExtraTreesClassifier(
            n_estimators=400,
            max_depth=12,
            min_samples_leaf=2,
            class_weight="balanced",
            random_state=RANDOM_STATE,
            n_jobs=-1,
        ),
        "Gradient Boosting": GradientBoostingClassifier(
            n_estimators=250,
            learning_rate=0.04,
            max_depth=3,
            min_samples_leaf=5,
            random_state=RANDOM_STATE,
        ),
        "Random Forest": RandomForestClassifier(
            n_estimators=400,
            max_depth=12,
            min_samples_leaf=2,
            class_weight="balanced_subsample",
            random_state=RANDOM_STATE,
            n_jobs=-1,
        ),
        "SVM": SVC(
            kernel="rbf",
            C=2.0,
            gamma="scale",
            probability=True,
            class_weight="balanced",
            random_state=RANDOM_STATE,
        ),
        "KNN": KNeighborsClassifier(
            n_neighbors=7,
            weights="distance",
        ),
    }

    if XGBOOST_AVAILABLE:
        models["XGBoost"] = XGBClassifier(
            objective="multi:softprob",
            num_class=len(CLASS_LABELS),
            n_estimators=350,
            learning_rate=0.04,
            max_depth=5,
            subsample=0.85,
            colsample_bytree=0.85,
            reg_alpha=0.5,
            reg_lambda=1.5,
            eval_metric="mlogloss",
            random_state=RANDOM_STATE,
            n_jobs=-1,
        )

    return models


def make_pipeline(model: Any) -> Pipeline:
    return Pipeline(
        steps=[
            ("imputer", SimpleImputer(strategy="median")),
            ("scaler", StandardScaler()),
            ("model", clone(model)),
        ]
    )


def supports_sample_weight(model_name: str) -> bool:
    return model_name in {
        "Gradient Boosting",
        "XGBoost",
    }


def fit_pipeline(
    pipeline: Pipeline,
    model_name: str,
    X: pd.DataFrame,
    y_int: pd.Series,
) -> Pipeline:
    if supports_sample_weight(model_name):
        weights = compute_sample_weight(
            class_weight="balanced",
            y=y_int,
        )
        pipeline.fit(X, y_int, model__sample_weight=weights)
    else:
        pipeline.fit(X, y_int)
    return pipeline


def decode_predictions(values: Iterable[int]) -> np.ndarray:
    return np.array(
        [INT_TO_CLASS[int(value)] for value in values],
        dtype=object,
    )


def overall_metrics(
    y_true: Iterable[str],
    y_pred: Iterable[str],
) -> dict[str, Any]:
    y_true_series = pd.Series(list(y_true), dtype="object")
    y_pred_series = pd.Series(list(y_pred), dtype="object")

    present = [
        label
        for label in CLASS_LABELS
        if int((y_true_series == label).sum()) > 0
    ]

    accuracy = accuracy_score(y_true_series, y_pred_series)
    balanced = balanced_accuracy_score(y_true_series, y_pred_series)

    macro_precision = precision_score(
        y_true_series,
        y_pred_series,
        labels=present,
        average="macro",
        zero_division=0,
    )
    macro_recall = recall_score(
        y_true_series,
        y_pred_series,
        labels=present,
        average="macro",
        zero_division=0,
    )
    macro_f1 = f1_score(
        y_true_series,
        y_pred_series,
        labels=present,
        average="macro",
        zero_division=0,
    )
    weighted_f1 = f1_score(
        y_true_series,
        y_pred_series,
        labels=present,
        average="weighted",
        zero_division=0,
    )
    mcc = (
        np.nan
        if y_true_series.nunique() < 2
        else float(matthews_corrcoef(y_true_series, y_pred_series))
    )

    drop_support = int((y_true_series == "Drop").sum())
    rise_support = int((y_true_series == "Rise").sum())

    drop_recall = (
        recall_score(
            y_true_series,
            y_pred_series,
            labels=["Drop"],
            average="macro",
            zero_division=0,
        )
        if drop_support
        else np.nan
    )
    rise_recall = (
        recall_score(
            y_true_series,
            y_pred_series,
            labels=["Rise"],
            average="macro",
            zero_division=0,
        )
        if rise_support
        else np.nan
    )

    drop_fn = int(
        ((y_true_series == "Drop") & (y_pred_series != "Drop")).sum()
    )
    rise_fn = int(
        ((y_true_series == "Rise") & (y_pred_series != "Rise")).sum()
    )
    drop_tp = int(
        ((y_true_series == "Drop") & (y_pred_series == "Drop")).sum()
    )
    rise_tp = int(
        ((y_true_series == "Rise") & (y_pred_series == "Rise")).sum()
    )

    return {
        "Accuracy": float(accuracy),
        "BalancedAccuracy": float(balanced),
        "MacroPrecision": float(macro_precision),
        "MacroRecall": float(macro_recall),
        "MacroF1": float(macro_f1),
        "WeightedF1": float(weighted_f1),
        "MCC": mcc,
        "DropRecall": (
            float(drop_recall) if not pd.isna(drop_recall) else np.nan
        ),
        "RiseRecall": (
            float(rise_recall) if not pd.isna(rise_recall) else np.nan
        ),
        "DropTruePositives": drop_tp,
        "DropFalseNegatives": drop_fn,
        "DropSupport": drop_support,
        "RiseTruePositives": rise_tp,
        "RiseFalseNegatives": rise_fn,
        "RiseSupport": rise_support,
        "ClassesPresent": ", ".join(present),
    }


def classwise_metrics(
    y_true: Iterable[str],
    y_pred: Iterable[str],
    split_name: str,
    model_name: str,
) -> pd.DataFrame:
    precision, recall, f1, support = precision_recall_fscore_support(
        list(y_true),
        list(y_pred),
        labels=CLASS_LABELS,
        zero_division=0,
    )

    rows = []
    for index, label in enumerate(CLASS_LABELS):
        rows.append(
            {
                "Split": split_name,
                "Model": model_name,
                "Class": label,
                "Precision": float(precision[index]),
                "Recall": float(recall[index]),
                "F1": float(f1[index]),
                "Support": int(support[index]),
            }
        )
    return pd.DataFrame(rows)


def date_based_cv_splits(
    train_df: pd.DataFrame,
    n_splits: int,
) -> list[tuple[np.ndarray, np.ndarray]]:
    unique_dates = np.array(sorted(train_df[DATE_COL].unique()))
    if len(unique_dates) < n_splits + 2:
        return []

    fold_size = len(unique_dates) // (n_splits + 1)
    if fold_size < 1:
        return []

    splits = []
    for fold in range(n_splits):
        train_date_end = fold_size * (fold + 1)
        validation_date_end = (
            len(unique_dates)
            if fold == n_splits - 1
            else fold_size * (fold + 2)
        )

        train_dates = unique_dates[:train_date_end]
        validation_dates = unique_dates[
            train_date_end:validation_date_end
        ]

        train_index = np.where(
            train_df[DATE_COL].isin(train_dates).to_numpy()
        )[0]
        validation_index = np.where(
            train_df[DATE_COL].isin(validation_dates).to_numpy()
        )[0]

        if len(train_index) and len(validation_index):
            splits.append((train_index, validation_index))

    return splits


def run_cross_validation(
    train_df: pd.DataFrame,
    feature_columns: list[str],
    n_splits: int,
) -> pd.DataFrame:
    rows = []
    splits = date_based_cv_splits(train_df, n_splits)
    if not splits:
        return pd.DataFrame()

    X_all = train_df[feature_columns]
    y_all = train_df[TARGET_COL].map(CLASS_TO_INT).astype(int)

    for model_name, model in get_models().items():
        for fold_number, (train_index, validation_index) in enumerate(
            splits,
            start=1,
        ):
            X_train = X_all.iloc[train_index]
            y_train = y_all.iloc[train_index]
            X_validation = X_all.iloc[validation_index]
            y_validation_text = train_df.iloc[
                validation_index
            ][TARGET_COL]

            if y_train.nunique() < 2:
                continue

            pipeline = make_pipeline(model)
            pipeline = fit_pipeline(
                pipeline,
                model_name,
                X_train,
                y_train,
            )
            predictions = decode_predictions(
                pipeline.predict(X_validation)
            )
            metrics = overall_metrics(
                y_validation_text,
                predictions,
            )

            rows.append(
                {
                    "Model": model_name,
                    "Fold": fold_number,
                    "TrainRows": len(train_index),
                    "ValidationRows": len(validation_index),
                    **metrics,
                }
            )

    return pd.DataFrame(rows)


def save_confusion_matrix(
    y_true: Iterable[str],
    y_pred: Iterable[str],
    output_path: Path,
    normalize: bool,
) -> None:
    matrix = confusion_matrix(
        list(y_true),
        list(y_pred),
        labels=CLASS_LABELS,
        normalize="true" if normalize else None,
    )

    fig, ax = plt.subplots(figsize=(6.8, 5.8))
    image = ax.imshow(
        matrix,
        vmin=0,
        vmax=1 if normalize else None,
    )
    fig.colorbar(image, ax=ax, fraction=0.046, pad=0.04)

    ax.set_xticks(range(len(CLASS_LABELS)))
    ax.set_xticklabels(CLASS_LABELS, fontsize=11)
    ax.set_yticks(range(len(CLASS_LABELS)))
    ax.set_yticklabels(CLASS_LABELS, fontsize=11)
    ax.set_xlabel("Predicted class", fontsize=12)
    ax.set_ylabel("True class", fontsize=12)

    for row in range(matrix.shape[0]):
        for column in range(matrix.shape[1]):
            value = matrix[row, column]
            text = f"{value:.2f}" if normalize else f"{int(value)}"
            ax.text(
                column,
                row,
                text,
                ha="center",
                va="center",
                fontsize=12,
            )

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_binary_confusion_matrix(
    y_true: Iterable[str],
    y_pred: Iterable[str],
    positive_class: str,
    output_path: Path,
) -> None:
    true_binary = (
        pd.Series(list(y_true), dtype="object") == positive_class
    ).astype(int)
    pred_binary = (
        pd.Series(list(y_pred), dtype="object") == positive_class
    ).astype(int)

    matrix = confusion_matrix(
        true_binary,
        pred_binary,
        labels=[0, 1],
    )

    fig, ax = plt.subplots(figsize=(6.2, 5.4))
    image = ax.imshow(matrix)
    fig.colorbar(image, ax=ax, fraction=0.046, pad=0.04)

    negative_name = f"Not {positive_class}"
    labels = [negative_name, positive_class]
    ax.set_xticks([0, 1])
    ax.set_xticklabels(labels, fontsize=11)
    ax.set_yticks([0, 1])
    ax.set_yticklabels(labels, fontsize=11)
    ax.set_xlabel("Predicted class", fontsize=12)
    ax.set_ylabel("True class", fontsize=12)

    for row in range(2):
        for column in range(2):
            ax.text(
                column,
                row,
                f"{int(matrix[row, column])}",
                ha="center",
                va="center",
                fontsize=13,
            )

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_model_comparison(
    metrics_df: pd.DataFrame,
    output_path: Path,
) -> None:
    subset = metrics_df[
        (metrics_df["Split"] == "Test")
        & (metrics_df["Status"] == "OK")
    ].copy()
    if subset.empty:
        return

    metrics = [
        "Accuracy",
        "BalancedAccuracy",
        "MacroF1",
        "MCC",
    ]
    x = np.arange(len(subset))
    width = 0.18

    fig, ax = plt.subplots(figsize=(11, 6.2))
    for metric_index, metric in enumerate(metrics):
        values = subset[metric].fillna(0.0).to_numpy()
        bars = ax.bar(
            x + (metric_index - 1.5) * width,
            values,
            width,
            label=metric,
        )
        for bar, value in zip(bars, values):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height(),
                f"{value:.2f}",
                ha="center",
                va="bottom",
                fontsize=8,
                rotation=90,
            )

    ax.set_xticks(x)
    ax.set_xticklabels(subset["Model"], rotation=20)
    ax.set_ylabel("Score", fontsize=12)
    ax.set_ylim(-0.05, 1.08)
    ax.legend(frameon=False, ncol=2)
    ax.grid(axis="y", alpha=0.3)

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_accuracy_balance_comparison(
    metrics_df: pd.DataFrame,
    output_path: Path,
) -> None:
    subset = metrics_df[
        (metrics_df["Split"] == "Test")
        & (metrics_df["Status"] == "OK")
    ].copy()
    if subset.empty:
        return

    x = np.arange(len(subset))
    width = 0.36

    fig, ax = plt.subplots(figsize=(10, 5.8))
    accuracy_bars = ax.bar(
        x - width / 2,
        subset["Accuracy"],
        width,
        label="Accuracy",
    )
    balanced_bars = ax.bar(
        x + width / 2,
        subset["BalancedAccuracy"],
        width,
        label="Balanced Accuracy",
    )

    for bars in [accuracy_bars, balanced_bars]:
        for bar in bars:
            value = bar.get_height()
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                value,
                f"{value:.3f}",
                ha="center",
                va="bottom",
                fontsize=9,
            )

    ax.set_xticks(x)
    ax.set_xticklabels(subset["Model"], rotation=20)
    ax.set_ylabel("Score", fontsize=12)
    ax.set_ylim(0, 1.08)
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=0.3)

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_class_distribution(
    target_df: pd.DataFrame,
    split: DataSplit,
    output_path: Path,
) -> None:
    records = []
    for name, frame in [
        ("All", target_df),
        ("Train", split.train),
        ("Validation", split.validation),
        ("Test", split.test),
    ]:
        counts = frame[TARGET_COL].value_counts()
        for label in CLASS_LABELS:
            records.append(
                {
                    "Split": name,
                    "Class": label,
                    "Count": int(counts.get(label, 0)),
                }
            )

    chart_df = pd.DataFrame(records)
    pivot = chart_df.pivot(
        index="Split",
        columns="Class",
        values="Count",
    ).reindex(["All", "Train", "Validation", "Test"])

    fig, ax = plt.subplots(figsize=(8.5, 5.8))
    bottom = np.zeros(len(pivot))
    x = np.arange(len(pivot))

    for label in CLASS_LABELS:
        values = pivot[label].to_numpy()
        ax.bar(
            x,
            values,
            bottom=bottom,
            label=label,
        )
        bottom += values

    ax.set_xticks(x)
    ax.set_xticklabels(pivot.index)
    ax.set_ylabel("Sample count", fontsize=12)
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=0.3)

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_horizon_histogram(
    target_df: pd.DataFrame,
    output_path: Path,
) -> None:
    fig, ax = plt.subplots(figsize=(8.5, 5.6))
    ax.hist(
        target_df[HORIZON_COL].dropna().to_numpy(),
        bins=min(20, target_df[HORIZON_COL].nunique()),
    )
    ax.set_xlabel(
        "Days to next available observation",
        fontsize=12,
    )
    ax.set_ylabel("Observation pairs", fontsize=12)
    ax.grid(axis="y", alpha=0.3)

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_temporal_prediction_plot(
    prediction_df: pd.DataFrame,
    output_path: Path,
) -> None:
    temporary = prediction_df.copy()
    temporary["YearMonth"] = (
        pd.to_datetime(temporary[DATE_COL])
        .dt.to_period("M")
        .astype(str)
    )
    trend = (
        temporary.groupby(["YearMonth", "PredictedTrend"])
        .size()
        .unstack(fill_value=0)
    )
    for label in CLASS_LABELS:
        if label not in trend.columns:
            trend[label] = 0
    trend = trend[CLASS_LABELS]

    fig, ax = plt.subplots(figsize=(10, 5.8))
    x = np.arange(len(trend))
    for label in CLASS_LABELS:
        ax.plot(
            x,
            trend[label].to_numpy(),
            marker="o",
            label=label,
        )

    ax.set_xticks(x)
    ax.set_xticklabels(trend.index, rotation=30, ha="right")
    ax.set_xlabel("Month", fontsize=12)
    ax.set_ylabel("Predicted test labels", fontsize=12)
    ax.legend(frameon=False)
    ax.grid(alpha=0.3)

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_feature_importance(
    pipeline: Pipeline,
    feature_columns: list[str],
    output_path: Path,
) -> bool:
    model = pipeline.named_steps["model"]
    if not hasattr(model, "feature_importances_"):
        return False

    importances = pd.Series(
        model.feature_importances_,
        index=feature_columns,
    ).sort_values(ascending=False).head(15)

    fig, ax = plt.subplots(figsize=(9, 6.8))
    ax.barh(
        importances.index[::-1],
        importances.values[::-1],
    )
    ax.set_xlabel("Feature importance", fontsize=12)
    ax.grid(axis="x", alpha=0.3)

    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return True


def threshold_sensitivity(
    base_df: pd.DataFrame,
    thresholds: list[float],
) -> pd.DataFrame:
    rows = []
    grouped = base_df.groupby(LOCATION_ID_COL, group_keys=False)
    next_hf = grouped[HF_COL].shift(-1)
    delta = next_hf - base_df[HF_COL]
    valid = delta.notna()

    for threshold in thresholds:
        labels = np.select(
            [
                delta[valid] < -threshold,
                delta[valid] > threshold,
            ],
            ["Drop", "Rise"],
            default="Stable",
        )
        counts = pd.Series(labels).value_counts()
        total = int(valid.sum())
        rows.append(
            {
                "Threshold": threshold,
                "TotalLabeledPairs": total,
                "Drop": int(counts.get("Drop", 0)),
                "Stable": int(counts.get("Stable", 0)),
                "Rise": int(counts.get("Rise", 0)),
                "DropPercent": (
                    100.0 * counts.get("Drop", 0) / total
                    if total
                    else np.nan
                ),
                "StablePercent": (
                    100.0 * counts.get("Stable", 0) / total
                    if total
                    else np.nan
                ),
                "RisePercent": (
                    100.0 * counts.get("Rise", 0) / total
                    if total
                    else np.nan
                ),
            }
        )

    return pd.DataFrame(rows)


def horizon_summary_table(target_df: pd.DataFrame) -> pd.DataFrame:
    series = target_df[HORIZON_COL].dropna()
    return pd.DataFrame(
        [
            {
                "Metric": "Count",
                "Value": len(series),
            },
            {
                "Metric": "MinimumDays",
                "Value": series.min(),
            },
            {
                "Metric": "FirstQuartileDays",
                "Value": series.quantile(0.25),
            },
            {
                "Metric": "MedianDays",
                "Value": series.median(),
            },
            {
                "Metric": "MeanDays",
                "Value": series.mean(),
            },
            {
                "Metric": "ThirdQuartileDays",
                "Value": series.quantile(0.75),
            },
            {
                "Metric": "MaximumDays",
                "Value": series.max(),
            },
        ]
    )


def horizon_distribution_table(target_df: pd.DataFrame) -> pd.DataFrame:
    counts = (
        target_df[HORIZON_COL]
        .value_counts()
        .sort_index()
        .rename_axis("HorizonDays")
        .reset_index(name="Count")
    )
    counts["Percent"] = (
        counts["Count"] / counts["Count"].sum() * 100.0
    )
    return counts


def split_summary_table(split: DataSplit) -> pd.DataFrame:
    rows = []
    for name, frame in [
        ("Train", split.train),
        ("Validation", split.validation),
        ("Test", split.test),
    ]:
        counts = frame[TARGET_COL].value_counts()
        rows.append(
            {
                "Split": name,
                "Rows": len(frame),
                "UniqueDates": frame[DATE_COL].nunique(),
                "StartDate": frame[DATE_COL].min().date().isoformat(),
                "EndDate": frame[DATE_COL].max().date().isoformat(),
                "Drop": int(counts.get("Drop", 0)),
                "Stable": int(counts.get("Stable", 0)),
                "Rise": int(counts.get("Rise", 0)),
            }
        )
    return pd.DataFrame(rows)


def label_summary_table(
    target_df: pd.DataFrame,
    threshold: float,
) -> pd.DataFrame:
    counts = target_df[TARGET_COL].value_counts()
    total = len(target_df)
    rows = []
    for label in CLASS_LABELS:
        count = int(counts.get(label, 0))
        rows.append(
            {
                "Class": label,
                "Count": count,
                "Percent": (
                    100.0 * count / total if total else np.nan
                ),
                "Threshold": threshold,
                "Definition": (
                    f"{DELTA_COL} < -{threshold}"
                    if label == "Drop"
                    else f"-{threshold} <= {DELTA_COL} <= {threshold}"
                    if label == "Stable"
                    else f"{DELTA_COL} > {threshold}"
                ),
            }
        )
    return pd.DataFrame(rows)


def leakage_audit_table(
    feature_columns: list[str],
    split: DataSplit,
) -> pd.DataFrame:
    forbidden_features = [
        column
        for column in feature_columns
        if column in FORBIDDEN_FEATURE_TERMS
        or column.endswith("_score")
    ]

    return pd.DataFrame(
        [
            {
                "Check": "Future target columns excluded from model input",
                "Pass": not forbidden_features,
                "Evidence": ", ".join(forbidden_features) or "None",
            },
            {
                "Check": "Train and validation dates do not overlap",
                "Pass": not bool(
                    set(split.train_dates)
                    & set(split.validation_dates)
                ),
                "Evidence": "Unique-date split",
            },
            {
                "Check": "Train and test dates do not overlap",
                "Pass": not bool(
                    set(split.train_dates)
                    & set(split.test_dates)
                ),
                "Evidence": "Unique-date split",
            },
            {
                "Check": "Validation and test dates do not overlap",
                "Pass": not bool(
                    set(split.validation_dates)
                    & set(split.test_dates)
                ),
                "Evidence": "Unique-date split",
            },
            {
                "Check": "Rolling statistics use only prior observations",
                "Pass": True,
                "Evidence": "group.shift(1).rolling(...)",
            },
            {
                "Check": "Imputer and scaler fit only inside training pipeline",
                "Pass": True,
                "Evidence": "SimpleImputer and StandardScaler are Pipeline steps.",
            },
            {
                "Check": "Test metrics excluded from model selection",
                "Pass": True,
                "Evidence": "Selection uses validation MacroF1, MCC, minimum Drop/Rise recall, BalancedAccuracy.",
            },
            {
                "Check": "Last observation per location excluded from labels",
                "Pass": True,
                "Evidence": "Rows with missing NextHealthFactor/NextObservationDate are dropped.",
            },
        ]
    )


def run_models(
    split: DataSplit,
    feature_columns: list[str],
    output_dir: Path,
) -> tuple[
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
    str,
    Pipeline,
]:
    figures_dir = output_dir / "figures"
    figures_dir.mkdir(parents=True, exist_ok=True)

    metrics_rows = []
    classwise_frames = []
    prediction_frames = []

    best_model_name: str | None = None
    best_pipeline: Pipeline | None = None
    best_rank: tuple[float, float, float, float] | None = None

    X_train = split.train[feature_columns]
    y_train_int = split.train[TARGET_COL].map(
        CLASS_TO_INT
    ).astype(int)

    for model_name, model in get_models().items():
        pipeline = make_pipeline(model)
        pipeline = fit_pipeline(
            pipeline,
            model_name,
            X_train,
            y_train_int,
        )

        for split_name, frame in [
            ("Train", split.train),
            ("Validation", split.validation),
            ("Test", split.test),
        ]:
            predictions = decode_predictions(
                pipeline.predict(frame[feature_columns])
            )
            metrics = overall_metrics(
                frame[TARGET_COL],
                predictions,
            )

            metrics_rows.append(
                {
                    "Split": split_name,
                    "Model": model_name,
                    "Status": "OK",
                    "Rows": len(frame),
                    **metrics,
                }
            )

            classwise_frames.append(
                classwise_metrics(
                    frame[TARGET_COL],
                    predictions,
                    split_name,
                    model_name,
                )
            )

            if split_name == "Test":
                prediction_frame = frame[
                    [
                        DATE_COL,
                        LOCATION_ID_COL,
                        LOCATION_NAME_COL,
                        HF_COL,
                        NEXT_HF_COL,
                        DELTA_COL,
                        HORIZON_COL,
                        TARGET_COL,
                    ]
                ].copy()
                prediction_frame["Model"] = model_name
                prediction_frame["PredictedTrend"] = predictions

                if hasattr(pipeline, "predict_proba"):
                    probabilities = pipeline.predict_proba(
                        frame[feature_columns]
                    )
                    model_classes = pipeline.named_steps[
                        "model"
                    ].classes_
                    for class_index, integer_class in enumerate(
                        model_classes
                    ):
                        class_label = INT_TO_CLASS[int(integer_class)]
                        prediction_frame[
                            f"Probability_{class_label}"
                        ] = probabilities[:, class_index]

                prediction_frames.append(prediction_frame)

                safe_name = model_name.replace(" ", "_")
                save_confusion_matrix(
                    frame[TARGET_COL],
                    predictions,
                    figures_dir
                    / f"cm_counts_{safe_name}_test.png",
                    normalize=False,
                )
                save_confusion_matrix(
                    frame[TARGET_COL],
                    predictions,
                    figures_dir
                    / f"cm_normalized_{safe_name}_test.png",
                    normalize=True,
                )

            if split_name == "Validation":
                min_event_recall = min(
                    metrics["DropRecall"],
                    metrics["RiseRecall"],
                )
                rank = (
                    metrics["MacroF1"],
                    (
                        metrics["MCC"]
                        if not pd.isna(metrics["MCC"])
                        else -1.0
                    ),
                    min_event_recall,
                    metrics["BalancedAccuracy"],
                )
                if best_rank is None or rank > best_rank:
                    best_rank = rank
                    best_model_name = model_name
                    best_pipeline = pipeline

    if best_model_name is None or best_pipeline is None:
        raise RuntimeError("Validation üzerinden en iyi model seçilemedi.")

    metrics_df = pd.DataFrame(metrics_rows)
    classwise_df = pd.concat(
        classwise_frames,
        ignore_index=True,
    )
    predictions_df = pd.concat(
        prediction_frames,
        ignore_index=True,
    )

    return (
        metrics_df,
        classwise_df,
        predictions_df,
        best_model_name,
        best_pipeline,
    )


def legacy_result_audit(
    best_model_name: str,
    metrics_df: pd.DataFrame,
) -> pd.DataFrame:
    best_test = metrics_df[
        (metrics_df["Model"] == best_model_name)
        & (metrics_df["Split"] == "Test")
    ].iloc[0]

    current_drop_fraction = (
        f"{int(best_test['DropTruePositives'])}/"
        f"{int(best_test['DropSupport'])}"
    )
    current_rise_fraction = (
        f"{int(best_test['RiseTruePositives'])}/"
        f"{int(best_test['RiseSupport'])}"
    )

    return pd.DataFrame(
        [
            {
                "LegacyStatement": "54/74 correctly predicted Drop events",
                "CurrentOutput": current_drop_fraction,
                "CurrentMetric": best_test["DropRecall"],
                "Status": (
                    "MATCH"
                    if current_drop_fraction == "54/74"
                    else "NOT_REPRODUCED_WITH_CURRENT_DATA"
                ),
            },
            {
                "LegacyStatement": "59/74 correctly predicted Drop events",
                "CurrentOutput": current_drop_fraction,
                "CurrentMetric": best_test["DropRecall"],
                "Status": (
                    "MATCH"
                    if current_drop_fraction == "59/74"
                    else "NOT_REPRODUCED_WITH_CURRENT_DATA"
                ),
            },
            {
                "LegacyStatement": "Approximately 80% Drop recall",
                "CurrentOutput": current_drop_fraction,
                "CurrentMetric": best_test["DropRecall"],
                "Status": "USE_CURRENT_CODE_OUTPUT",
            },
            {
                "LegacyStatement": "51/60 correctly predicted Rise events",
                "CurrentOutput": current_rise_fraction,
                "CurrentMetric": best_test["RiseRecall"],
                "Status": (
                    "MATCH"
                    if current_rise_fraction == "51/60"
                    else "NOT_REPRODUCED_WITH_CURRENT_DATA"
                ),
            },
        ]
    )


def figure_manifest_table(
    best_model_name: str,
    feature_importance_created: bool,
) -> pd.DataFrame:
    safe_name = best_model_name.replace(" ", "_")
    rows = [
        {
            "ManuscriptFigure": "Figure 7",
            "File": f"figures/figure_07_best_multiclass_confusion_matrix.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": f"Test confusion matrix for selected model ({best_model_name}).",
        },
        {
            "ManuscriptFigure": "Figure 8",
            "File": "figures/figure_08_drop_binary_confusion_matrix.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": "One-vs-rest Drop confusion matrix for selected model.",
        },
        {
            "ManuscriptFigure": "Figure 9",
            "File": "figures/figure_09_rise_binary_confusion_matrix.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": "One-vs-rest Rise confusion matrix for selected model.",
        },
        {
            "ManuscriptFigure": "Figure 10",
            "File": "figures/figure_10_accuracy_vs_balanced_accuracy.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": "Accuracy and balanced accuracy comparison across models.",
        },
        {
            "ManuscriptFigure": "Figure 11",
            "File": "figures/figure_11_temporal_predicted_labels.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": "Monthly temporal distribution of selected model predictions on test data.",
        },
        {
            "ManuscriptFigure": "Supplementary",
            "File": "figures/proactive_model_comparison.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": "Accuracy, balanced accuracy, Macro F1 and MCC comparison.",
        },
        {
            "ManuscriptFigure": "Supplementary",
            "File": "figures/trend_class_distribution.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": "Drop, Stable and Rise distribution by split.",
        },
        {
            "ManuscriptFigure": "Supplementary",
            "File": "figures/horizon_days_distribution.png",
            "GeneratedBy": "proactive_trend_model.py",
            "Description": "Actual days to the next available observation.",
        },
    ]

    if feature_importance_created:
        rows.append(
            {
                "ManuscriptFigure": "Supplementary",
                "File": "figures/best_model_feature_importance.png",
                "GeneratedBy": "proactive_trend_model.py",
                "Description": f"Top feature importances for {best_model_name}.",
            }
        )

    return pd.DataFrame(rows)


def _safe_number(value: Any, digits: int = 4) -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):.{digits}f}"


def write_word_report(
    output_path: Path,
    input_path: Path,
    threshold: float,
    target_df: pd.DataFrame,
    feature_columns: list[str],
    split_summary: pd.DataFrame,
    label_summary: pd.DataFrame,
    horizon_summary: pd.DataFrame,
    metrics_df: pd.DataFrame,
    classwise_df: pd.DataFrame,
    cv_df: pd.DataFrame,
    leakage_df: pd.DataFrame,
    legacy_audit_df: pd.DataFrame,
    figure_manifest_df: pd.DataFrame,
    best_model_name: str,
) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.shared import Cm, Pt
    except ImportError:
        print(
            "[UYARI] python-docx kurulu değil; Word raporu üretilemedi."
        )
        return

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9.5)

    def format_table(table, font_size: float = 8.5) -> None:
        """Keep rows intact, repeat the header and use a compact readable font."""
        for row_index, row in enumerate(table.rows):
            row_properties = row._tr.get_or_add_trPr()

            cannot_split = OxmlElement("w:cantSplit")
            row_properties.append(cannot_split)

            if row_index == 0:
                repeat_header = OxmlElement("w:tblHeader")
                repeat_header.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
                row_properties.append(repeat_header)

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(font_size)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Proactive Layer Technical Report")
    run.bold = True
    run.font.size = Pt(17)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Next-available-observation Drop / Stable / Rise classification"
    ).italic = True

    document.add_heading("1. Scope and terminology", level=1)
    document.add_paragraph(
        f"Input file: {input_path}. The Proactive Layer classifies the "
        "direction of change between the current HF observation and the next "
        "available HF observation for the same physical sampling point. It "
        "does not predict a fixed calendar day or a guaranteed one-week "
        "future value. Therefore, the scientifically accurate term is "
        "'next-observation trend classification'."
    )

    document.add_heading("2. Target-label generation", level=1)
    document.add_paragraph(
        f"Delta_HF(i,t) = HF(i,t+1) - HF(i,t). Drop is assigned when "
        f"Delta_HF < -{threshold}; Stable when -{threshold} <= Delta_HF <= "
        f"{threshold}; and Rise when Delta_HF > {threshold}. The t+1 record "
        "is obtained with a group-wise shift(-1) operation using NoktaId. "
        "The final record of every sampling point has no future label and is "
        "excluded from supervised modeling."
    )

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    columns = [
        "Class",
        "Count",
        "Percent",
        "Threshold",
        "Definition",
    ]
    for index, column in enumerate(columns):
        table.rows[0].cells[index].text = column
    for _, row in label_summary.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["Class"])
        cells[1].text = str(row["Count"])
        cells[2].text = _safe_number(row["Percent"], 2)
        cells[3].text = str(row["Threshold"])
        cells[4].text = str(row["Definition"])
    format_table(table, 8.5)

    document.add_heading("3. Actual prediction horizon", level=1)
    horizon_map = dict(
        zip(horizon_summary["Metric"], horizon_summary["Value"])
    )
    document.add_paragraph(
        "The actual time to the next observation is irregular. In the current "
        f"data, the median horizon is {_safe_number(horizon_map.get('MedianDays'), 1)} "
        f"days, the mean is {_safe_number(horizon_map.get('MeanDays'), 1)} days, "
        f"and the observed range is {_safe_number(horizon_map.get('MinimumDays'), 0)} "
        f"to {_safe_number(horizon_map.get('MaximumDays'), 0)} days. HorizonDays "
        "is reported for interpretation but is not used as a predictor because "
        "the next measurement date is not necessarily known at prediction time."
    )

    document.add_heading("4. Features and leakage prevention", level=1)
    document.add_paragraph(
        "The predictors consist of current raw water-quality measurements, "
        "current HF, lagged HF values, past-only rolling statistics, the "
        "previous measurement gap, calendar seasonality, and coordinates "
        "when available. NextHealthFactor, FutureDeltaHF, HorizonDays, "
        "TrendLabel, RiskClass, WAWQI, FailFast and *_score columns are "
        "excluded from the predictor matrix."
    )
    document.add_paragraph(
        "Feature list: " + ", ".join(feature_columns)
    )

    audit_table = document.add_table(rows=1, cols=3)
    audit_table.style = "Table Grid"
    for index, header in enumerate(["Audit check", "Pass", "Evidence"]):
        audit_table.rows[0].cells[index].text = header
    for _, row in leakage_df.iterrows():
        cells = audit_table.add_row().cells
        cells[0].text = str(row["Check"])
        cells[1].text = "PASS" if bool(row["Pass"]) else "FAIL"
        cells[2].text = str(row["Evidence"])
    format_table(audit_table, 8.2)

    document.add_heading("5. Time-aware data split", level=1)
    split_table = document.add_table(rows=1, cols=8)
    split_table.style = "Table Grid"
    split_columns = [
        "Split",
        "Rows",
        "UniqueDates",
        "StartDate",
        "EndDate",
        "Drop",
        "Stable",
        "Rise",
    ]
    for index, column in enumerate(split_columns):
        split_table.rows[0].cells[index].text = column
    for _, row in split_summary.iterrows():
        cells = split_table.add_row().cells
        for index, column in enumerate(split_columns):
            cells[index].text = str(row[column])
    format_table(split_table, 8.0)

    document.add_paragraph(
        "The split is performed by unique observation dates. No date appears "
        "in more than one split. Missing-value imputation and scaling are "
        "fitted inside each model pipeline using training data only. Model "
        "selection is based on validation performance; the test set is used "
        "only for final evaluation."
    )

    document.add_heading("6. Compared models and selection rule", level=1)
    document.add_paragraph(
        "The following models are compared: Extra Trees, Gradient Boosting, "
        "XGBoost, Random Forest, SVM and KNN. XGBoost is treated as a distinct "
        "algorithm and is never described as Extra Trees. The selected model "
        "is determined lexicographically by validation Macro F1, MCC, the "
        "minimum of Drop Recall and Rise Recall, and Balanced Accuracy. "
        f"The selected model is {best_model_name}."
    )

    document.add_heading("7. Final test performance", level=1)
    test_metrics = metrics_df[
        (metrics_df["Split"] == "Test")
        & (metrics_df["Status"] == "OK")
    ].copy()

    performance_columns = [
        "Model",
        "Accuracy",
        "BalancedAccuracy",
        "MacroPrecision",
        "MacroRecall",
        "MacroF1",
        "MCC",
        "DropRecall",
        "RiseRecall",
        "DropFalseNegatives",
        "RiseFalseNegatives",
    ]
    table = document.add_table(
        rows=1,
        cols=len(performance_columns),
    )
    table.style = "Table Grid"
    performance_headers = [
        "Model",
        "Accuracy",
        "BalAcc",
        "MacroPrec",
        "MacroRec",
        "MacroF1",
        "MCC",
        "DropRec",
        "RiseRec",
        "DropFN",
        "RiseFN",
    ]
    for index, header in enumerate(performance_headers):
        table.rows[0].cells[index].text = header

    for _, row in test_metrics.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(performance_columns):
            value = row[column]
            if column in {
                "DropFalseNegatives",
                "RiseFalseNegatives",
            }:
                cells[index].text = str(int(value))
            elif column == "Model":
                cells[index].text = str(value)
            else:
                cells[index].text = _safe_number(value)
    format_table(table, 7.2)

    document.add_heading("8. Class-wise metrics for selected model", level=1)
    best_classwise = classwise_df[
        (classwise_df["Split"] == "Test")
        & (classwise_df["Model"] == best_model_name)
    ]
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    class_columns = [
        "Class",
        "Precision",
        "Recall",
        "F1",
        "Support",
        "Model",
    ]
    for index, column in enumerate(class_columns):
        table.rows[0].cells[index].text = column
    for _, row in best_classwise.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(class_columns):
            value = row[column]
            if column in {"Precision", "Recall", "F1"}:
                cells[index].text = _safe_number(value)
            else:
                cells[index].text = str(value)
    format_table(table, 8.3)

    document.add_heading("9. Expanding-window cross-validation", level=1)
    if cv_df.empty:
        document.add_paragraph(
            "Cross-validation could not be calculated because the available "
            "date sequence was insufficient."
        )
    else:
        cv_summary = (
            cv_df.groupby("Model")[
                [
                    "BalancedAccuracy",
                    "MacroF1",
                    "MCC",
                    "DropRecall",
                    "RiseRecall",
                ]
            ]
            .agg(["mean", "std"])
            .reset_index()
        )
        document.add_paragraph(
            "Expanding date-based folds are calculated using training dates "
            "only. Fold-level values are exported in proactive_cv_metrics.csv."
        )
        for _, row in cv_summary.iterrows():
            model_name = row[("Model", "")]
            document.add_paragraph(
                f"{model_name}: mean Macro F1 "
                f"{_safe_number(row[('MacroF1', 'mean')])}, mean MCC "
                f"{_safe_number(row[('MCC', 'mean')])}, mean Drop Recall "
                f"{_safe_number(row[('DropRecall', 'mean')])}, mean Rise Recall "
                f"{_safe_number(row[('RiseRecall', 'mean')])}.",
                style="List Bullet",
            )

    document.add_heading("10. Legacy result consistency audit", level=1)
    document.add_paragraph(
        "The manuscript previously included conflicting Drop counts such as "
        "54/74 and 59/74. The current data, target-generation rule and selected "
        "model produce one reproducible numerator and denominator. Legacy "
        "numbers that do not match the current output must be removed or "
        "explicitly identified as results from an obsolete experiment."
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    audit_columns = [
        "LegacyStatement",
        "CurrentOutput",
        "CurrentMetric",
        "Status",
    ]
    for index, column in enumerate(audit_columns):
        table.rows[0].cells[index].text = column
    for _, row in legacy_audit_df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["LegacyStatement"])
        cells[1].text = str(row["CurrentOutput"])
        cells[2].text = _safe_number(row["CurrentMetric"])
        status_text = {
            "NOT_REPRODUCED_WITH_CURRENT_DATA": "Not reproduced",
            "USE_CURRENT_CODE_OUTPUT": "Use current output",
            "MATCH": "Match",
        }.get(str(row["Status"]), str(row["Status"]))
        cells[3].text = status_text
    format_table(table, 8.2)

    document.add_heading("11. Figure mapping", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    figure_columns = [
        "ManuscriptFigure",
        "File",
        "GeneratedBy",
        "Description",
    ]
    for index, column in enumerate(figure_columns):
        table.rows[0].cells[index].text = column
    for _, row in figure_manifest_df.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(figure_columns):
            cells[index].text = str(row[column])
    format_table(table, 7.8)

    document.add_heading("12. Interpretation limitation", level=1)
    document.add_paragraph(
        "This pipeline predicts the direction of the next observed HF value. "
        "It does not constitute classical fixed-step time-series forecasting. "
        "The horizon varies according to the actual measurement schedule. "
        "Therefore, claims about a daily or weekly forecast must not be made "
        "unless a separate fixed-calendar forecasting design is implemented."
    )

    document.save(output_path)


def run_pipeline(args: argparse.Namespace) -> None:
    input_path = find_input_csv(args.input)
    output_dir = Path(args.output_dir)
    figures_dir = output_dir / "figures"
    output_dir.mkdir(parents=True, exist_ok=True)
    figures_dir.mkdir(parents=True, exist_ok=True)

    print(f"[LOAD] {input_path}")
    raw_df = load_data(input_path)

    sensitivity_df = threshold_sensitivity(
        raw_df,
        thresholds=sorted(
            set(
                [
                    args.threshold,
                    0.05,
                    0.10,
                    0.20,
                    0.30,
                    0.50,
                ]
            )
        ),
    )

    target_df = create_trend_target(
        raw_df,
        threshold=args.threshold,
    )
    feature_df, feature_columns = create_past_only_features(
        target_df
    )

    split = split_by_unique_dates(
        feature_df,
        train_ratio=args.train_ratio,
        validation_ratio=args.validation_ratio,
    )

    print(
        f"[DATA] labeled_rows={len(feature_df)} "
        f"dates={feature_df[DATE_COL].nunique()} "
        f"locations={feature_df[LOCATION_ID_COL].nunique()}"
    )
    print("[TREND DISTRIBUTION]")
    print(feature_df[TARGET_COL].value_counts())

    cv_df = run_cross_validation(
        split.train,
        feature_columns,
        n_splits=args.cv_splits,
    )

    (
        metrics_df,
        classwise_df,
        predictions_df,
        best_model_name,
        best_pipeline,
    ) = run_models(
        split,
        feature_columns,
        output_dir,
    )

    split_summary_df = split_summary_table(split)
    label_summary_df = label_summary_table(
        feature_df,
        args.threshold,
    )
    horizon_summary_df = horizon_summary_table(feature_df)
    horizon_distribution_df = horizon_distribution_table(
        feature_df
    )
    leakage_df = leakage_audit_table(
        feature_columns,
        split,
    )
    legacy_audit_df = legacy_result_audit(
        best_model_name,
        metrics_df,
    )

    best_test_predictions = predictions_df[
        predictions_df["Model"] == best_model_name
    ].copy()

    save_model_comparison(
        metrics_df,
        figures_dir / "proactive_model_comparison.png",
    )
    save_accuracy_balance_comparison(
        metrics_df,
        figures_dir
        / "figure_10_accuracy_vs_balanced_accuracy.png",
    )
    save_class_distribution(
        feature_df,
        split,
        figures_dir / "trend_class_distribution.png",
    )
    save_horizon_histogram(
        feature_df,
        figures_dir / "horizon_days_distribution.png",
    )
    save_temporal_prediction_plot(
        best_test_predictions,
        figures_dir
        / "figure_11_temporal_predicted_labels.png",
    )

    save_confusion_matrix(
        best_test_predictions[TARGET_COL],
        best_test_predictions["PredictedTrend"],
        figures_dir
        / "figure_07_best_multiclass_confusion_matrix.png",
        normalize=False,
    )
    save_binary_confusion_matrix(
        best_test_predictions[TARGET_COL],
        best_test_predictions["PredictedTrend"],
        "Drop",
        figures_dir
        / "figure_08_drop_binary_confusion_matrix.png",
    )
    save_binary_confusion_matrix(
        best_test_predictions[TARGET_COL],
        best_test_predictions["PredictedTrend"],
        "Rise",
        figures_dir
        / "figure_09_rise_binary_confusion_matrix.png",
    )

    feature_importance_created = save_feature_importance(
        best_pipeline,
        feature_columns,
        figures_dir / "best_model_feature_importance.png",
    )

    figure_manifest_df = figure_manifest_table(
        best_model_name,
        feature_importance_created,
    )

    metrics_df.to_csv(
        output_dir / "proactive_model_metrics.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    classwise_df.to_csv(
        output_dir / "proactive_classwise_metrics.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    predictions_df.to_csv(
        output_dir / "proactive_predictions.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    cv_df.to_csv(
        output_dir / "proactive_cv_metrics.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    split_summary_df.to_csv(
        output_dir / "proactive_split_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )
    label_summary_df.to_csv(
        output_dir / "trend_label_summary.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.4f",
    )
    horizon_summary_df.to_csv(
        output_dir / "trend_horizon_summary.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.4f",
    )
    horizon_distribution_df.to_csv(
        output_dir / "trend_horizon_distribution.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.4f",
    )
    sensitivity_df.to_csv(
        output_dir / "trend_threshold_sensitivity.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.4f",
    )
    pd.DataFrame(
        {
            "Feature": feature_columns,
            "Definition": [
                (
                    "Current raw or current HF feature"
                    if column
                    in RAW_PARAMETER_CANDIDATES + ["HF_Current"]
                    else "Past-only engineered or calendar feature"
                )
                for column in feature_columns
            ],
        }
    ).to_csv(
        output_dir / "proactive_feature_list.csv",
        index=False,
        encoding="utf-8-sig",
    )
    leakage_df.to_csv(
        output_dir / "data_leakage_audit.csv",
        index=False,
        encoding="utf-8-sig",
    )
    legacy_audit_df.to_csv(
        output_dir / "legacy_result_audit.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    figure_manifest_df.to_csv(
        output_dir / "figure_manifest.csv",
        index=False,
        encoding="utf-8-sig",
    )

    model_bundle = {
        "pipeline": best_pipeline,
        "model_name": best_model_name,
        "feature_columns": feature_columns,
        "class_labels": CLASS_LABELS,
        "class_to_int": CLASS_TO_INT,
        "int_to_class": INT_TO_CLASS,
        "threshold": args.threshold,
        "target_definition": (
            "Drop if NextHF-CurrentHF < -threshold; "
            "Stable if within threshold; Rise if > threshold."
        ),
        "forecast_horizon": "Next available observation per NoktaId",
        "input_file": str(input_path),
        "selection_rule": (
            "Validation ranking by MacroF1, MCC, minimum of "
            "DropRecall/RiseRecall, then BalancedAccuracy"
        ),
        "training_date_range": [
            min(split.train_dates).date().isoformat(),
            max(split.train_dates).date().isoformat(),
        ],
    }
    joblib.dump(
        model_bundle,
        output_dir / "best_proactive_model.joblib",
    )

    write_word_report(
        output_dir / "Proactive_Layer_Technical_Report.docx",
        input_path,
        args.threshold,
        feature_df,
        feature_columns,
        split_summary_df,
        label_summary_df,
        horizon_summary_df,
        metrics_df,
        classwise_df,
        cv_df,
        leakage_df,
        legacy_audit_df,
        figure_manifest_df,
        best_model_name,
    )

    print("\n[SUCCESS] Proactive Layer pipeline completed.")
    print(f"[BEST MODEL] {best_model_name}")
    print(f"[OUTPUT] {output_dir}")
    print("\n[TEST METRICS]")
    print(
        metrics_df[
            metrics_df["Split"] == "Test"
        ][
            [
                "Model",
                "Accuracy",
                "BalancedAccuracy",
                "MacroF1",
                "MCC",
                "DropRecall",
                "RiseRecall",
                "DropTruePositives",
                "DropSupport",
                "RiseTruePositives",
                "RiseSupport",
            ]
        ].to_string(index=False)
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Run reviewer-ready Proactive Layer trend classification."
        )
    )
    parser.add_argument(
        "--input",
        default=None,
        help="Path to izsu_features.csv",
    )
    parser.add_argument(
        "--output-dir",
        default="outputs/proactive",
    )
    parser.add_argument(
        "--threshold",
        type=float,
        default=0.05,
        help=(
            "Absolute HF tolerance. Default preserves the previous "
            "code rule: Drop < -0.05, Rise > 0.05."
        ),
    )
    parser.add_argument(
        "--train-ratio",
        type=float,
        default=0.70,
    )
    parser.add_argument(
        "--validation-ratio",
        type=float,
        default=0.15,
    )
    parser.add_argument(
        "--cv-splits",
        type=int,
        default=4,
    )
    return parser.parse_args()


if __name__ == "__main__":
    run_pipeline(parse_args())
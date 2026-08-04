
"""
hf_next_observation_regression.py

Predicts the Health Factor at the next AVAILABLE observation for the same
sampling point. It does not claim a fixed one-week forecast.

Leakage prevention
------------------
- Target is generated with group-wise shift(-1) by NoktaId.
- Last observation of each location is removed from supervised training.
- All lag/rolling features use current or prior information only.
- Train/validation/test split is chronological by unique dates.
- Imputer and scaler are fitted inside the training pipeline.
- Model selection uses validation MAE/RMSE/R2, never test performance.

Outputs include model metrics, persistence baseline, predictions, per-location
next-observation estimates, 300 DPI figures and an automatic Word report.
"""

from __future__ import annotations

import argparse
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import joblib
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from sklearn.base import clone
from sklearn.ensemble import (
    GradientBoostingRegressor,
    RandomForestRegressor,
)
from sklearn.impute import SimpleImputer
from sklearn.linear_model import Ridge
from sklearn.metrics import (
    mean_absolute_error,
    mean_squared_error,
    median_absolute_error,
    r2_score,
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import RobustScaler

warnings.filterwarnings("ignore")

try:
    from xgboost import XGBRegressor

    XGBOOST_AVAILABLE = True
except Exception:
    XGBOOST_AVAILABLE = False


DATE_COL = "Tarih"
LOCATION_ID_COL = "NoktaId"
LOCATION_NAME_COL = "NoktaAdi"
HF_COL = "HealthFactor"
TARGET_COL = "NextHealthFactor"
NEXT_DATE_COL = "NextObservationDate"
HORIZON_COL = "HorizonDays"
RANDOM_STATE = 42

RAW_PARAMETER_CANDIDATES = [
    "Alüminyum",
    "Amonyum",
    "Arsenik",
    "Bulanıklık",
    "C.Perfringens",
    "Demir",
    "E.coli",
    "Klorür",
    "Koku",
    "Koliform Bakteri",
    "Nitrit",
    "Oksitlenebilirlik",
    "Renk",
    "Toplam Sertlik",
    "Tuzluluk",
    "pH",
    "İletkenlik",
]


@dataclass(frozen=True)
class RegressionSplit:
    train: pd.DataFrame
    validation: pd.DataFrame
    test: pd.DataFrame
    train_dates: list[pd.Timestamp]
    validation_dates: list[pd.Timestamp]
    test_dates: list[pd.Timestamp]


def find_input_csv(requested: str | None) -> Path:
    if requested:
        path = Path(requested)
        if path.exists():
            return path
        raise FileNotFoundError(f"Girdi CSV bulunamadı: {path}")

    script_dir = Path(__file__).resolve().parent
    candidates = [
        Path("data/data/izsu_features.csv"),
        Path("data/izsu_features.csv"),
        Path("izsu_features.csv"),
        script_dir.parent / "data" / "data" / "izsu_features.csv",
        script_dir.parent / "data" / "izsu_features.csv",
    ]
    for path in candidates:
        if path.exists():
            return path

    raise FileNotFoundError(
        "izsu_features.csv bulunamadı. --input ile dosya yolunu belirtin."
    )


def load_data(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, encoding="utf-8-sig")
    required = {DATE_COL, LOCATION_NAME_COL, HF_COL}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Eksik zorunlu sütunlar: {sorted(missing)}")

    if LOCATION_ID_COL not in df.columns:
        df[LOCATION_ID_COL] = df[LOCATION_NAME_COL].astype(str)

    df[DATE_COL] = pd.to_datetime(df[DATE_COL], errors="coerce")
    df[HF_COL] = pd.to_numeric(df[HF_COL], errors="coerce")
    df = df.dropna(
        subset=[DATE_COL, LOCATION_ID_COL, LOCATION_NAME_COL, HF_COL]
    ).copy()

    if "E.Coli" in df.columns and df["E.Coli"].isna().all():
        df = df.drop(columns=["E.Coli"])

    for column in RAW_PARAMETER_CANDIDATES + ["Enlem", "Boylam"]:
        if column in df.columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")

    duplicate_count = int(
        df.duplicated(subset=[DATE_COL, LOCATION_ID_COL]).sum()
    )
    if duplicate_count:
        raise ValueError(
            f"Tarih + NoktaId anahtarında {duplicate_count} duplicate var."
        )

    return df.sort_values(
        [LOCATION_ID_COL, DATE_COL], kind="stable"
    ).reset_index(drop=True)


def create_features(
    raw_df: pd.DataFrame,
) -> tuple[pd.DataFrame, list[str], pd.DataFrame]:
    df = raw_df.copy()
    group = df.groupby(LOCATION_ID_COL, group_keys=False)

    df[TARGET_COL] = group[HF_COL].shift(-1)
    df[NEXT_DATE_COL] = group[DATE_COL].shift(-1)
    df[HORIZON_COL] = (df[NEXT_DATE_COL] - df[DATE_COL]).dt.days

    df["HF_Current"] = df[HF_COL]
    df["HF_Lag1"] = group[HF_COL].shift(1)
    df["HF_Lag2"] = group[HF_COL].shift(2)
    df["HF_Lag3"] = group[HF_COL].shift(3)

    df["HF_PastMean3"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(3, min_periods=1).mean()
    )
    df["HF_PastStd3"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(3, min_periods=2).std()
    )
    df["HF_PastMean7"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(7, min_periods=1).mean()
    )
    df["HF_PastStd7"] = group[HF_COL].transform(
        lambda series: series.shift(1).rolling(7, min_periods=2).std()
    )

    df["HF_ChangeFromPrevious"] = df["HF_Current"] - df["HF_Lag1"]
    df["HF_PreviousChange"] = df["HF_Lag1"] - df["HF_Lag2"]
    df["HF_DifferenceFromPastMean3"] = (
        df["HF_Current"] - df["HF_PastMean3"]
    )

    previous_date = group[DATE_COL].shift(1)
    df["PreviousGapDays"] = (df[DATE_COL] - previous_date).dt.days
    df["MonthSin"] = np.sin(2 * np.pi * df[DATE_COL].dt.month / 12.0)
    df["MonthCos"] = np.cos(2 * np.pi * df[DATE_COL].dt.month / 12.0)

    feature_columns = [
        column
        for column in RAW_PARAMETER_CANDIDATES
        if column in df.columns and not df[column].isna().all()
    ]
    feature_columns += [
        "HF_Current",
        "HF_Lag1",
        "HF_Lag2",
        "HF_Lag3",
        "HF_PastMean3",
        "HF_PastStd3",
        "HF_PastMean7",
        "HF_PastStd7",
        "HF_ChangeFromPrevious",
        "HF_PreviousChange",
        "HF_DifferenceFromPastMean3",
        "PreviousGapDays",
        "MonthSin",
        "MonthCos",
    ]
    for coordinate in ["Enlem", "Boylam"]:
        if coordinate in df.columns and not df[coordinate].isna().all():
            feature_columns.append(coordinate)

    feature_columns = list(dict.fromkeys(feature_columns))

    inference_df = df.copy()
    supervised_df = df.dropna(
        subset=[TARGET_COL, NEXT_DATE_COL, HORIZON_COL]
    ).copy()

    if int((supervised_df[HORIZON_COL] <= 0).sum()):
        raise ValueError("Pozitif olmayan HorizonDays tespit edildi.")

    return (
        supervised_df.reset_index(drop=True),
        feature_columns,
        inference_df.reset_index(drop=True),
    )


def split_by_dates(
    df: pd.DataFrame,
    train_ratio: float,
    validation_ratio: float,
) -> RegressionSplit:
    dates = sorted(pd.Series(df[DATE_COL].unique()).tolist())
    if len(dates) < 8:
        raise ValueError("Zamansal ayrım için yeterli benzersiz tarih yok.")

    train_end = max(1, int(len(dates) * train_ratio))
    validation_end = max(
        train_end + 1,
        int(len(dates) * (train_ratio + validation_ratio)),
    )
    validation_end = min(validation_end, len(dates) - 1)

    train_dates = dates[:train_end]
    validation_dates = dates[train_end:validation_end]
    test_dates = dates[validation_end:]

    train = df[df[DATE_COL].isin(train_dates)].copy()
    validation = df[df[DATE_COL].isin(validation_dates)].copy()
    test = df[df[DATE_COL].isin(test_dates)].copy()

    if train.empty or validation.empty or test.empty:
        raise ValueError("Train/validation/test bölümlerinden biri boş.")

    return RegressionSplit(
        train=train,
        validation=validation,
        test=test,
        train_dates=train_dates,
        validation_dates=validation_dates,
        test_dates=test_dates,
    )


def get_models() -> dict[str, Any]:
    models: dict[str, Any] = {
        "Ridge": Ridge(alpha=1.0),
        "Random Forest": RandomForestRegressor(
            n_estimators=400,
            max_depth=10,
            min_samples_leaf=3,
            random_state=RANDOM_STATE,
            n_jobs=-1,
        ),
        "Gradient Boosting": GradientBoostingRegressor(
            n_estimators=350,
            learning_rate=0.03,
            max_depth=3,
            min_samples_leaf=5,
            subsample=0.85,
            random_state=RANDOM_STATE,
        ),
    }

    if XGBOOST_AVAILABLE:
        models["XGBoost"] = XGBRegressor(
            objective="reg:squarederror",
            n_estimators=500,
            learning_rate=0.03,
            max_depth=5,
            subsample=0.85,
            colsample_bytree=0.85,
            reg_alpha=0.5,
            reg_lambda=1.5,
            random_state=RANDOM_STATE,
            n_jobs=-1,
            verbosity=0,
        )

    return models


def make_pipeline(model: Any) -> Pipeline:
    return Pipeline(
        steps=[
            ("imputer", SimpleImputer(strategy="median")),
            ("scaler", RobustScaler()),
            ("model", clone(model)),
        ]
    )


def safe_mape(
    y_true: np.ndarray,
    y_pred: np.ndarray,
    epsilon: float = 1e-8,
) -> float:
    denominator = np.maximum(np.abs(y_true), epsilon)
    return float(np.mean(np.abs((y_true - y_pred) / denominator)))


def evaluate(
    y_true: pd.Series | np.ndarray,
    y_pred: np.ndarray,
) -> dict[str, float]:
    true_array = np.asarray(y_true, dtype=float)
    pred_array = np.asarray(y_pred, dtype=float)

    return {
        "R2": float(r2_score(true_array, pred_array)),
        "RMSE": float(
            np.sqrt(mean_squared_error(true_array, pred_array))
        ),
        "MAE": float(mean_absolute_error(true_array, pred_array)),
        "MedAE": float(
            median_absolute_error(true_array, pred_array)
        ),
        "MAPE": safe_mape(true_array, pred_array),
    }


def run_models(
    split: RegressionSplit,
    feature_columns: list[str],
) -> tuple[pd.DataFrame, pd.DataFrame, str, Pipeline]:
    rows = []
    prediction_frames = []

    best_name: str | None = None
    best_pipeline: Pipeline | None = None
    best_rank: tuple[float, float, float] | None = None

    for model_name, model in get_models().items():
        pipeline = make_pipeline(model)
        pipeline.fit(
            split.train[feature_columns],
            split.train[TARGET_COL],
        )

        for split_name, frame in [
            ("Train", split.train),
            ("Validation", split.validation),
            ("Test", split.test),
        ]:
            predictions = pipeline.predict(frame[feature_columns])
            metrics = evaluate(frame[TARGET_COL], predictions)
            rows.append(
                {
                    "Split": split_name,
                    "Model": model_name,
                    "Rows": len(frame),
                    **metrics,
                }
            )

            if split_name == "Test":
                prediction_frame = frame[
                    [
                        DATE_COL,
                        LOCATION_ID_COL,
                        LOCATION_NAME_COL,
                        HF_COL,
                        TARGET_COL,
                        NEXT_DATE_COL,
                        HORIZON_COL,
                    ]
                ].copy()
                prediction_frame["Model"] = model_name
                prediction_frame["PredictedNextHealthFactor"] = predictions
                prediction_frame["Residual"] = (
                    prediction_frame[TARGET_COL]
                    - prediction_frame["PredictedNextHealthFactor"]
                )
                prediction_frames.append(prediction_frame)

            if split_name == "Validation":
                rank = (
                    -metrics["MAE"],
                    -metrics["RMSE"],
                    metrics["R2"],
                )
                if best_rank is None or rank > best_rank:
                    best_rank = rank
                    best_name = model_name
                    best_pipeline = pipeline

    if best_name is None or best_pipeline is None:
        raise RuntimeError("Validation üzerinden en iyi model seçilemedi.")

    return (
        pd.DataFrame(rows),
        pd.concat(prediction_frames, ignore_index=True),
        best_name,
        best_pipeline,
    )


def baseline_metrics(split: RegressionSplit) -> pd.DataFrame:
    rows = []
    for split_name, frame in [
        ("Train", split.train),
        ("Validation", split.validation),
        ("Test", split.test),
    ]:
        predictions = frame[HF_COL].to_numpy()
        rows.append(
            {
                "Split": split_name,
                "Model": "Persistence Baseline",
                "Rows": len(frame),
                **evaluate(frame[TARGET_COL], predictions),
            }
        )
    return pd.DataFrame(rows)


def split_summary(split: RegressionSplit) -> pd.DataFrame:
    rows = []
    for name, frame in [
        ("Train", split.train),
        ("Validation", split.validation),
        ("Test", split.test),
    ]:
        rows.append(
            {
                "Split": name,
                "Rows": len(frame),
                "UniqueDates": frame[DATE_COL].nunique(),
                "StartDate": frame[DATE_COL].min().date().isoformat(),
                "EndDate": frame[DATE_COL].max().date().isoformat(),
                "MeanTargetHF": frame[TARGET_COL].mean(),
                "StdTargetHF": frame[TARGET_COL].std(),
            }
        )
    return pd.DataFrame(rows)


def horizon_summary(df: pd.DataFrame) -> pd.DataFrame:
    series = df[HORIZON_COL].dropna()
    return pd.DataFrame(
        [
            {"Metric": "Count", "Value": len(series)},
            {"Metric": "MinimumDays", "Value": series.min()},
            {
                "Metric": "FirstQuartileDays",
                "Value": series.quantile(0.25),
            },
            {"Metric": "MedianDays", "Value": series.median()},
            {"Metric": "MeanDays", "Value": series.mean()},
            {
                "Metric": "ThirdQuartileDays",
                "Value": series.quantile(0.75),
            },
            {"Metric": "MaximumDays", "Value": series.max()},
        ]
    )


def leakage_audit(
    feature_columns: list[str],
    split: RegressionSplit,
) -> pd.DataFrame:
    forbidden = {
        TARGET_COL,
        NEXT_DATE_COL,
        HORIZON_COL,
        "FutureDeltaHF",
        "TrendLabel",
        "RiskClass",
        "WAWQI",
        "FailFast",
    }
    leaks = [
        column
        for column in feature_columns
        if column in forbidden or column.endswith("_score")
    ]

    return pd.DataFrame(
        [
            {
                "Check": "Future target fields excluded from X",
                "Pass": not leaks,
                "Evidence": ", ".join(leaks) or "None",
            },
            {
                "Check": "Train/validation dates do not overlap",
                "Pass": not bool(
                    set(split.train_dates)
                    & set(split.validation_dates)
                ),
                "Evidence": "Unique-date split",
            },
            {
                "Check": "Train/test dates do not overlap",
                "Pass": not bool(
                    set(split.train_dates) & set(split.test_dates)
                ),
                "Evidence": "Unique-date split",
            },
            {
                "Check": "Validation/test dates do not overlap",
                "Pass": not bool(
                    set(split.validation_dates)
                    & set(split.test_dates)
                ),
                "Evidence": "Unique-date split",
            },
            {
                "Check": "Rolling statistics use past observations only",
                "Pass": True,
                "Evidence": "group.shift(1).rolling(...)",
            },
            {
                "Check": "Imputer/scaler fitted through training pipeline",
                "Pass": True,
                "Evidence": "SimpleImputer + RobustScaler inside Pipeline",
            },
            {
                "Check": "Best model selected without test metrics",
                "Pass": True,
                "Evidence": "Validation MAE, RMSE and R2",
            },
        ]
    )


def create_latest_forecasts(
    inference_df: pd.DataFrame,
    feature_columns: list[str],
    pipeline: Pipeline,
    model_name: str,
) -> pd.DataFrame:
    latest = (
        inference_df.sort_values(
            [LOCATION_ID_COL, DATE_COL],
            kind="stable",
        )
        .groupby(LOCATION_ID_COL, as_index=False)
        .tail(1)
        .copy()
    )
    latest["PredictedNextHealthFactor"] = pipeline.predict(
        latest[feature_columns]
    )
    latest["BestModel"] = model_name
    latest["ForecastMeaning"] = "Next available observation; date unknown"

    columns = [
        DATE_COL,
        LOCATION_ID_COL,
        LOCATION_NAME_COL,
        HF_COL,
        "PredictedNextHealthFactor",
        "BestModel",
        "ForecastMeaning",
    ]
    return latest[columns].sort_values(
        "PredictedNextHealthFactor",
        ascending=False,
    )


def save_actual_vs_predicted(
    prediction_df: pd.DataFrame,
    output_path: Path,
) -> None:
    fig, ax = plt.subplots(figsize=(7, 6))
    ax.scatter(
        prediction_df[TARGET_COL],
        prediction_df["PredictedNextHealthFactor"],
        alpha=0.7,
    )
    lower = min(
        prediction_df[TARGET_COL].min(),
        prediction_df["PredictedNextHealthFactor"].min(),
    )
    upper = max(
        prediction_df[TARGET_COL].max(),
        prediction_df["PredictedNextHealthFactor"].max(),
    )
    ax.plot([lower, upper], [lower, upper], linestyle="--")
    ax.set_xlabel("Actual next-observation HF", fontsize=12)
    ax.set_ylabel("Predicted next-observation HF", fontsize=12)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_residual_histogram(
    prediction_df: pd.DataFrame,
    output_path: Path,
) -> None:
    fig, ax = plt.subplots(figsize=(7.5, 5.5))
    ax.hist(prediction_df["Residual"], bins=30)
    ax.axvline(
        prediction_df["Residual"].mean(),
        linestyle="--",
    )
    ax.set_xlabel("Residual: actual - predicted", fontsize=12)
    ax.set_ylabel("Count", fontsize=12)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_residual_vs_predicted(
    prediction_df: pd.DataFrame,
    output_path: Path,
) -> None:
    fig, ax = plt.subplots(figsize=(7.5, 5.5))
    ax.scatter(
        prediction_df["PredictedNextHealthFactor"],
        prediction_df["Residual"],
        alpha=0.7,
    )
    ax.axhline(0.0, linestyle="--")
    ax.set_xlabel("Predicted next-observation HF", fontsize=12)
    ax.set_ylabel("Residual", fontsize=12)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_model_metrics_plot(
    metrics_df: pd.DataFrame,
    output_path: Path,
) -> None:
    test = metrics_df[metrics_df["Split"] == "Test"].copy()
    x = np.arange(len(test))
    width = 0.28

    mae_bars = test["MAE"].to_numpy()
    rmse_bars = test["RMSE"].to_numpy()

    fig, ax = plt.subplots(figsize=(9, 5.8))
    bars_mae = ax.bar(
        x - width / 2,
        mae_bars,
        width,
        label="MAE",
    )
    bars_rmse = ax.bar(
        x + width / 2,
        rmse_bars,
        width,
        label="RMSE",
    )

    for bars in [bars_mae, bars_rmse]:
        for bar in bars:
            value = bar.get_height()
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                value,
                f"{value:.3f}",
                ha="center",
                va="bottom",
                fontsize=9,
            )

    ax.set_xticks(x)
    ax.set_xticklabels(test["Model"], rotation=20)
    ax.set_ylabel("Error in HF units", fontsize=12)
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def _safe_number(value: Any, digits: int = 4) -> str:
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):.{digits}f}"


def write_report(
    output_path: Path,
    input_path: Path,
    feature_columns: list[str],
    split_df: pd.DataFrame,
    horizon_df: pd.DataFrame,
    metrics_df: pd.DataFrame,
    baseline_df: pd.DataFrame,
    audit_df: pd.DataFrame,
    best_model_name: str,
) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.shared import Cm, Pt
    except ImportError:
        print(
            "[UYARI] python-docx kurulu değil; regresyon Word raporu üretilemedi."
        )
        return

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9.5)

    def format_table(table, font_size: float = 8.3) -> None:
        for row_index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                header = OxmlElement("w:tblHeader")
                header.set(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val",
                    "true",
                )
                properties.append(header)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(font_size)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Next-Observation HF Regression Report")
    run.bold = True
    run.font.size = Pt(17)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Leakage-safe numerical prediction of the next available HF observation"
    ).italic = True

    document.add_heading("1. Scope", level=1)
    document.add_paragraph(
        f"Input file: {input_path}. This regression model estimates the HF "
        "value at the next available observation for the same NoktaId. It does "
        "not predict a fixed one-week date. The future observation interval is "
        "irregular and is documented separately."
    )

    document.add_heading("2. Target and horizon", level=1)
    document.add_paragraph(
        "The target is created with NextHealthFactor = groupby(NoktaId)"
        "[HealthFactor].shift(-1). The final observation of each sampling point "
        "has no numeric target and is excluded from supervised training."
    )
    horizon_map = dict(zip(horizon_df["Metric"], horizon_df["Value"]))
    document.add_paragraph(
        f"The observed horizon has a median of "
        f"{_safe_number(horizon_map.get('MedianDays'), 1)} days, a mean of "
        f"{_safe_number(horizon_map.get('MeanDays'), 1)} days and a range of "
        f"{_safe_number(horizon_map.get('MinimumDays'), 0)} to "
        f"{_safe_number(horizon_map.get('MaximumDays'), 0)} days."
    )

    document.add_heading("3. Predictors and leakage prevention", level=1)
    document.add_paragraph(
        "The input matrix contains current raw measurements, current HF, "
        "past HF lags and rolling statistics, previous measurement gap, "
        "calendar variables and coordinates. NextHealthFactor, future date, "
        "HorizonDays, trend labels, RiskClass, WAWQI, FailFast and *_score "
        "columns are excluded."
    )
    document.add_paragraph("Features: " + ", ".join(feature_columns))

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for index, value in enumerate(["Audit check", "Pass", "Evidence"]):
        table.rows[0].cells[index].text = value
    for _, row in audit_df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["Check"])
        cells[1].text = "PASS" if bool(row["Pass"]) else "FAIL"
        cells[2].text = str(row["Evidence"])
    format_table(table)

    document.add_heading("4. Time-aware data split", level=1)
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    columns = [
        "Split",
        "Rows",
        "UniqueDates",
        "StartDate",
        "EndDate",
        "MeanTargetHF",
        "StdTargetHF",
    ]
    for index, column in enumerate(columns):
        table.rows[0].cells[index].text = column
    for _, row in split_df.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            value = row[column]
            cells[index].text = (
                _safe_number(value)
                if column in {"MeanTargetHF", "StdTargetHF"}
                else str(value)
            )
    format_table(table, 7.8)

    document.add_heading("5. Model selection", level=1)
    document.add_paragraph(
        "Ridge, Random Forest, Gradient Boosting and XGBoost are compared. "
        "The best model is selected using validation MAE first, validation "
        "RMSE second and validation R2 third. Test metrics are not used for "
        f"selection. Selected model: {best_model_name}."
    )

    document.add_heading("6. Test performance and persistence baseline", level=1)
    combined = pd.concat([metrics_df, baseline_df], ignore_index=True)
    test = combined[combined["Split"] == "Test"].copy()
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    columns = ["Model", "R2", "RMSE", "MAE", "MedAE", "MAPE", "Rows"]
    for index, column in enumerate(columns):
        table.rows[0].cells[index].text = column
    for _, row in test.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            value = row[column]
            cells[index].text = (
                str(value)
                if column in {"Model", "Rows"}
                else _safe_number(value)
            )
    format_table(table, 8.0)

    document.add_paragraph(
        "The persistence baseline predicts that the next HF will equal the "
        "current HF. Because the series is highly stable, a machine-learning "
        "model should be interpreted relative to this baseline rather than "
        "from R2 alone."
    )

    document.add_heading("7. Interpretation limitation", level=1)
    document.add_paragraph(
        "The generated per-location forecast file represents the model's "
        "estimate for the next measurement whenever it occurs. A date is not "
        "assigned because the measurement schedule is not fixed. These results "
        "are supplementary to the main Reactive and Proactive classification "
        "layers."
    )

    document.save(output_path)


def run_pipeline(args: argparse.Namespace) -> None:
    input_path = find_input_csv(args.input)
    output_dir = Path(args.output_dir)
    figures_dir = output_dir / "figures"
    output_dir.mkdir(parents=True, exist_ok=True)
    figures_dir.mkdir(parents=True, exist_ok=True)

    print(f"[LOAD] {input_path}")
    raw_df = load_data(input_path)
    supervised_df, feature_columns, inference_df = create_features(raw_df)
    split = split_by_dates(
        supervised_df,
        args.train_ratio,
        args.validation_ratio,
    )

    metrics_df, predictions_df, best_name, best_pipeline = run_models(
        split,
        feature_columns,
    )
    baseline_df = baseline_metrics(split)
    split_df = split_summary(split)
    horizon_df = horizon_summary(supervised_df)
    audit_df = leakage_audit(feature_columns, split)

    best_predictions = predictions_df[
        predictions_df["Model"] == best_name
    ].copy()
    latest_forecasts = create_latest_forecasts(
        inference_df,
        feature_columns,
        best_pipeline,
        best_name,
    )

    metrics_df.to_csv(
        output_dir / "regression_model_metrics.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    baseline_df.to_csv(
        output_dir / "regression_baseline_metrics.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    predictions_df.to_csv(
        output_dir / "regression_predictions.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    latest_forecasts.to_csv(
        output_dir / "next_observation_forecasts.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    split_df.to_csv(
        output_dir / "regression_split_summary.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    horizon_df.to_csv(
        output_dir / "regression_horizon_summary.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )
    audit_df.to_csv(
        output_dir / "data_leakage_audit.csv",
        index=False,
        encoding="utf-8-sig",
    )
    pd.DataFrame({"Feature": feature_columns}).to_csv(
        output_dir / "regression_feature_list.csv",
        index=False,
        encoding="utf-8-sig",
    )

    save_actual_vs_predicted(
        best_predictions,
        figures_dir / "actual_vs_predicted_test.png",
    )
    save_residual_histogram(
        best_predictions,
        figures_dir / "residual_distribution_test.png",
    )
    save_residual_vs_predicted(
        best_predictions,
        figures_dir / "residual_vs_predicted_test.png",
    )
    save_model_metrics_plot(
        pd.concat([metrics_df, baseline_df], ignore_index=True),
        figures_dir / "regression_test_error_comparison.png",
    )

    bundle = {
        "pipeline": best_pipeline,
        "model_name": best_name,
        "feature_columns": feature_columns,
        "target": TARGET_COL,
        "forecast_horizon": "Next available observation per NoktaId",
        "selection_rule": "Validation MAE, then RMSE, then R2",
        "training_date_range": [
            min(split.train_dates).date().isoformat(),
            max(split.train_dates).date().isoformat(),
        ],
        "input_file": str(input_path),
    }
    joblib.dump(
        bundle,
        output_dir / "best_regression_model.joblib",
    )

    write_report(
        output_dir / "Regression_Technical_Report.docx",
        input_path,
        feature_columns,
        split_df,
        horizon_df,
        metrics_df,
        baseline_df,
        audit_df,
        best_name,
    )

    print("\n[SUCCESS] Regression pipeline completed.")
    print(f"[BEST MODEL] {best_name}")
    print(f"[OUTPUT] {output_dir}")
    print("\n[TEST METRICS]")
    print(
        pd.concat([metrics_df, baseline_df], ignore_index=True)
        .query("Split == 'Test'")[
            ["Model", "R2", "RMSE", "MAE", "MedAE", "MAPE"]
        ]
        .to_string(index=False)
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Next-available-observation HF regression."
    )
    parser.add_argument(
        "--input",
        default=None,
        help="Path to izsu_features.csv",
    )
    parser.add_argument(
        "--output-dir",
        default="outputs/regression",
    )
    parser.add_argument(
        "--train-ratio",
        type=float,
        default=0.70,
    )
    parser.add_argument(
        "--validation-ratio",
        type=float,
        default=0.15,
    )
    return parser.parse_args()


if __name__ == "__main__":
    run_pipeline(parse_args())
